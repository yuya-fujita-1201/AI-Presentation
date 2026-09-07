#!/usr/bin/env python3
"""Check the independently reviewable local delivery without changing deck sources."""
import argparse,hashlib,json,zipfile,importlib.util
from pathlib import Path
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
MANIFEST=ROOT/'docs/later-series-upgrade-2026-09-07/status.json'
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
 ap=argparse.ArgumentParser();ap.add_argument('--output',required=True);a=ap.parse_args()
 result=[]
 for row in json.loads(MANIFEST.read_text())['topics']:
  p=ROOT/row['deck'];d=json.loads((p/'deck.json').read_text());name=d['meta']['id'];ppt=p/'build'/f'{name}.pptx';html=p/'build'/f'{name}.html'
  sections=[str(s.get('number')) for s in d['slides'] if s['type']=='section']
  checks={'source_preserved':digest(ROOT/row['source']/'deck.json')==row['source_sha256'],'id_matches_folder':name==p.name,'four_chapters':all(f'{i:02}' in sections for i in range(1,5)),'html_exists':html.is_file(),'images_exist':all((p/s['path']).is_file() for s in d['slides'] if s.get('path')),'preview_count':len(list((p/'build/preview').glob('slide-*.png')))==len(d['slides'])}
  with zipfile.ZipFile(ppt) as z:checks['zip_valid']=z.testzip() is None
  pr=Presentation(ppt);checks['slide_count']=len(pr.slides)==len(d['slides']);checks['no_notes']=all(not s.has_notes_slide for s in pr.slides)
  r={'deck':row['deck'],'slides':len(d['slides']),'deck_sha256':digest(p/'deck.json'),'pptx_sha256':digest(ppt),'html_sha256':digest(html),'checks':checks,'passed':all(checks.values())};result.append(r)
 Path(a.output).write_text(json.dumps({'topics':result,'passed':all(r['passed'] for r in result)},ensure_ascii=False,indent=2)+'\n')
 print(json.dumps([{'deck':r['deck'],'slides':r['slides'],'passed':r['passed']} for r in result],ensure_ascii=False))
 raise SystemExit(0 if all(r['passed'] for r in result) else 1)
if __name__=='__main__':main()
