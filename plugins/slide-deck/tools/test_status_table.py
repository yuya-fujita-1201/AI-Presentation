#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""status_table（進捗ステータス表）のビルド・検証テスト。"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

TOOLS_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS_DIR))
import build_deck  # noqa: E402


def _slide(**kw):
    s = {"type": "status_table", "title": "進捗報告", "lead": "各チーム予定通り",
         "columns": ["チーム", "進捗状況", "実績", "計画"],
         "rows": [
             {"label": "PMO", "status": "green",
              "cells": [[{"text": "進捗管理", "children": ["週次レビュー", {"text": "論点", "children": ["対応方針"]}]}],
                        ["次週の計画"]]},
             {"label": "開発", "status": "yellow", "cells": [["実装"], ["テスト"]]}]}
    s.update(kw)
    return s


class FlattenTest(unittest.TestCase):
    def test_nested_items_become_levels(self):
        self.assertEqual([(0, "a"), (0, "x"), (1, "b"), (1, "c")],
                         build_deck.flatten_status_items(["a", {"text": "x", "children": ["b", "c"]}]))
        self.assertEqual([(0, "親"), (1, "子"), (2, "孫")],
                         build_deck.flatten_status_items([{"text": "親", "children": [{"text": "子", "children": ["孫"]}]}]))

    def test_status_color_resolves_legend_key_token_and_hex(self):
        s = _slide()
        self.assertEqual("#00B050", build_deck.status_color(s, "green"))
        self.assertEqual("accent", build_deck.status_color(s, "accent"))
        self.assertEqual("#123456", build_deck.status_color(s, "#123456"))


class BuildTest(unittest.TestCase):
    def _build(self, slide):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        deck_dir = Path(tmp.name) / "status"
        deck_dir.mkdir()
        (deck_dir / "deck.json").write_text(json.dumps(
            {"meta": {"title": "s"}, "slides": [slide]}, ensure_ascii=False), encoding="utf-8")
        proc = subprocess.run([sys.executable, str(TOOLS_DIR / "build_deck.py"), str(deck_dir)],
                              capture_output=True, text=True, encoding="utf-8", errors="replace")
        return deck_dir, proc

    def test_html_and_native_pptx_table(self):
        from pptx import Presentation
        deck_dir, proc = self._build(_slide())
        self.assertEqual(0, proc.returncode, proc.stderr)
        html = (deck_dir / "build" / "status.html").read_text(encoding="utf-8")
        for text in ("進捗管理", "週次レビュー", "対応方針", "✓", "●", "予定通りまたは前倒し"):
            self.assertIn(text, html)
        prs = Presentation(deck_dir / "build" / "status.pptx")
        self.assertFalse(prs.slides[0].has_notes_slide)
        tables = [sh.table for sh in prs.slides[0].shapes if sh.has_table]
        self.assertEqual(2, len(tables))  # 凡例表＋本体
        body = max(tables, key=lambda t: len(t.columns))  # 本体は4列、凡例は3列
        self.assertEqual(3, len(body.rows))      # 見出し＋2行
        self.assertEqual("PMO", body.cell(1, 0).text)
        self.assertEqual("●", body.cell(1, 1).text)
        self.assertIn("週次レビュー", body.cell(1, 2).text)

    def test_legend_false_hides_legend_table(self):
        from pptx import Presentation
        deck_dir, proc = self._build(_slide(legend=False))
        self.assertEqual(0, proc.returncode, proc.stderr)
        prs = Presentation(deck_dir / "build" / "status.pptx")
        self.assertEqual(1, sum(1 for sh in prs.slides[0].shapes if sh.has_table))

    def test_cells_count_mismatch_fails(self):
        s = _slide()
        s["rows"][0]["cells"] = [["実績だけ"]]
        _d, proc = self._build(s)
        self.assertNotEqual(0, proc.returncode)
        self.assertIn("cells 数", proc.stderr)

    def test_unknown_status_fails(self):
        s = _slide()
        s["rows"][0]["status"] = "blue-ish"
        _d, proc = self._build(s)
        self.assertNotEqual(0, proc.returncode)
        self.assertIn("status", proc.stderr)


if __name__ == "__main__":
    unittest.main()
