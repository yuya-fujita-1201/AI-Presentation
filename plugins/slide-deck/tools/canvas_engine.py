"""canvas 型スライドの描画エンジン（HTML / PPTX）。

canvas は 1280x720 の座標系に図形・文字・表・ベクターアイコンを座標指定で置くスライドタイプ。
既存 PPTX の見た目を崩さず取り込む用途（tools/pptx_import.py が生成する）を主目的とする。
要素スキーマは references/deck-schema.md の「canvas」節を参照。単位はすべて px（文字サイズも px）。
"""
from __future__ import annotations

import html as _html
import math
import re

EMU_PER_PX = 9525

GEOMS = {
    "rect", "roundRect", "ellipse", "triangle", "diamond", "homePlate", "chevron", "flowChartDecision",
    "flowChartMerge", "flowChartDocument", "can", "flowChartMagneticDisk", "borderCallout1",
    "borderCallout2", "line", "straightConnector1", "bentConnector2", "bentConnector3",
    "bentConnector4", "bentConnector5", "custom",
}
LINE_GEOMS = {"line", "straightConnector1", "bentConnector2", "bentConnector3", "bentConnector4", "bentConnector5"}
KINDS = {"shape", "table", "svg", "backdrop", "placeholder", "group"}

DASH = {"dash": (4, 3), "sysDash": (3, 1), "sysDot": (1, 1), "dot": (1, 3), "dashDot": (4, 3, 1, 3),
        "lgDash": (8, 3), "lgDashDot": (8, 3, 1, 3), "lgDashDotDot": (8, 3, 1, 3, 1, 3),
        "sysDashDot": (3, 1, 1, 1), "sysDashDotDot": (3, 1, 1, 1, 1, 1)}
ARROW_SZ = {"sm": 2.0, "med": 3.0, "lg": 5.0}
# PowerPoint の行間 100% に相当する line-height 倍率
LINE_HEIGHT = 1.2


def esc(s) -> str:
    return _html.escape(str(s), quote=True)


# ---------------------------------------------------------------- sanitize
# canvas の要素値は HTML の属性・CSS と PPTX の XML 属性へそのまま埋め込まれるため、描画の前に
# 必ず sanitize_elements() を通す（取り込み元の PPTX や手書きの deck.json は信用しない）。
# validate_canvas() も同じ関数で置き換えが起きた箇所を集めてエラーにする。
_HEX6_RE = re.compile(r"^#[0-9A-Fa-f]{6}$")
_WORD_RE = re.compile(r"^[A-Za-z0-9_]{1,40}$")
_FONT_BAD_RE = re.compile(r"[\"'<>&;{}\\\x00-\x1f]")
ENUMS = {
    "anchor": {"t", "ctr", "b"},
    "align": {"l", "ctr", "r", "just", "dist"},
    "cap": {"rnd", "sq", "flat"},
    "merged": {"h", "v"},
}
ARROW_TYPES = {"triangle", "stealth", "arrow", "oval", "diamond", "none"}
PATH_OPS = {"M", "L", "C", "Q", "Z"}
COLOR_KEYS = {"color", "highlight", "fg", "bg", "fill"}
_DROP = object()  # 不正な列挙値はキーごと落として既定値に任せる


def _num(v, default=0.0):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) else default


def _sanitize(v, key, path, issues):
    """要素ツリーを再帰的に複製しながら危険な値を安全な値へ置き換える。置き換えたら issues に記録"""
    def bad(msg, repl):
        issues.append(f"{path}: {msg}")
        return repl

    if key in ("span", "rowspan"):
        n = _num(v, None)
        return max(int(n), 1) if n is not None else bad(f"{key} '{v}' は整数で指定してください", 1)
    if key == "alpha":
        return min(max(_num(v, 1.0), 0.0), 1.0)
    if isinstance(v, dict):
        if key in ("head", "tail"):
            out = {}
            t = v.get("type", "triangle")
            out["type"] = t if t in ARROW_TYPES else bad(f"矢じり type '{t}' は未対応です", "triangle")
            for k in ("w", "len"):
                s = v.get(k, "med")
                out[k] = s if s in ARROW_SZ else bad(f"矢じり {k} '{s}' は sm/med/lg のいずれかです", "med")
            return out
        if key == "adj":
            return {k: _num(x) for k, x in v.items() if isinstance(k, str) and _WORD_RE.match(k)}
        out = {k: _sanitize(x, k, f"{path}.{k}", issues) for k, x in v.items()}
        return {k: x for k, x in out.items() if x is not _DROP}
    if isinstance(v, list):
        if key == "font":
            out = []
            for x in v:
                s = str(x)
                if _FONT_BAD_RE.search(s):
                    s = bad(f"フォント名 '{s}' に使えない文字が含まれています", _FONT_BAD_RE.sub("", s))
                out.append(s)
            return out
        if key == "cmds":
            out = []
            for c in v:
                if not (isinstance(c, list) and c and c[0] in PATH_OPS):
                    bad(f"パス命令 {c!r} は M/L/C/Q/Z のいずれかで始めてください", None)
                    continue
                out.append([c[0]] + [_num(x) for x in c[1:]])
            return out
        if key == "stops":
            return [[_num(s[0]), _sanitize(s[1], "color", f"{path}[{i}]", issues)]
                    for i, s in enumerate(v) if isinstance(s, (list, tuple)) and len(s) == 2]
        if key == "colors":
            return [_sanitize(c, "color", f"{path}[{i}]", issues) for i, c in enumerate(v)]
        return [_sanitize(x, key, f"{path}[{i}]", issues) for i, x in enumerate(v)]
    if isinstance(v, str):
        if key in COLOR_KEYS:
            return v if _HEX6_RE.match(v) else bad(f"色 '{v}' は #RRGGBB で指定してください", "#000000")
        if key in ENUMS and v not in ENUMS[key]:
            return bad(f"{key} '{v}' は {'/'.join(sorted(ENUMS[key]))} のいずれかです", _DROP)
        if key == "dash" and v not in DASH and v != "solid":
            return bad(f"dash '{v}' は未対応です", _DROP)
        if key == "pattern" and not _WORD_RE.match(v):
            return bad(f"pattern '{v}' は英数字のプリセット名で指定してください", "pct50")
        if key == "svg":
            clean = sanitize_svg(v)
            return clean if clean else bad("svg を解析できないか、許可されない内容です", "")
        return v
    return v


def sanitize_elements(elements, path="elements", issues=None):
    """canvas 要素列の安全な複製を返す（描画前に必ず通す）"""
    return _sanitize(elements if isinstance(elements, list) else [], "elements", path,
                     issues if issues is not None else [])


# SVG はベクターアイコン用途に限るため、要素・属性とも許可リスト方式で残す
_SVG_NS = "http://www.w3.org/2000/svg"
_XLINK_NS = "http://www.w3.org/1999/xlink"
_SVG_TAGS = {
    "svg", "g", "path", "rect", "circle", "ellipse", "line", "polyline", "polygon", "text", "tspan",
    "defs", "linearGradient", "radialGradient", "stop", "clipPath", "mask", "use", "symbol", "title",
    "desc", "pattern", "marker", "filter", "feGaussianBlur", "feOffset", "feBlend", "feFlood",
    "feComposite", "feColorMatrix", "feMerge", "feMergeNode",
}
_SVG_ATTRS = {
    "id", "class", "viewBox", "width", "height", "x", "y", "x1", "y1", "x2", "y2", "cx", "cy", "r",
    "rx", "ry", "fx", "fy", "d", "points", "transform", "fill", "fill-opacity", "fill-rule", "stroke",
    "stroke-width", "stroke-opacity", "stroke-linecap", "stroke-linejoin", "stroke-dasharray",
    "stroke-dashoffset", "stroke-miterlimit", "opacity", "clip-path", "clip-rule", "mask", "offset",
    "stop-color", "stop-opacity", "gradientUnits", "gradientTransform", "spreadMethod",
    "patternUnits", "patternContentUnits", "patternTransform", "clipPathUnits", "maskUnits",
    "maskContentUnits", "preserveAspectRatio", "font-family", "font-size", "font-weight",
    "font-style", "text-anchor", "dominant-baseline", "letter-spacing", "dx", "dy", "style",
    "visibility", "display", "markerWidth", "markerHeight", "refX", "refY", "orient", "markerUnits",
    "stdDeviation", "in", "in2", "result", "mode", "operator", "k1", "k2", "k3", "k4", "values",
    "type", "flood-color", "flood-opacity", "filterUnits", "primitiveUnits", "href", "version",
}
_SVG_BAD_STYLE_RE = re.compile(r"url\(\s*['\"]?\s*[^#'\"\s)]|javascript:|expression\s*\(|@import|behavior\s*:|[<>]", re.I)


def sanitize_svg(src) -> str:
    """インライン SVG を許可リストで浄化して返す。解析できない・DOCTYPE を含む場合は空文字"""
    import xml.etree.ElementTree as ET
    if not isinstance(src, str) or len(src) > 2_000_000 or re.search(r"<!(DOCTYPE|ENTITY)", src, re.I):
        return ""
    try:
        root = ET.fromstring(src)
    except ET.ParseError:
        return ""

    def local(name):
        return name.rsplit("}", 1)[-1]

    def clean(el):
        for child in list(el):
            if not isinstance(child.tag, str) or local(child.tag) not in _SVG_TAGS:
                el.remove(child)
            else:
                clean(child)
        for name in list(el.attrib):
            ln, val = local(name), el.attrib[name]
            ok = ln in _SVG_ATTRS
            if ok and ln == "href":
                ok = val.strip().startswith("#")
            elif ok and ln == "style":
                ok = not _SVG_BAD_STYLE_RE.search(val)
            elif ok and "url(" in val:
                ok = not _SVG_BAD_STYLE_RE.search(val)
            if not ok:
                del el.attrib[name]
            elif ln == "href" and name != "href":
                # xlink:href は SVG2 の href に寄せる（名前空間接頭辞なしで HTML へ直挿しできるように）
                del el.attrib[name]
                el.set("href", val)

    if local(root.tag) != "svg":
        return ""
    clean(root)
    for el in root.iter():
        if isinstance(el.tag, str):
            el.tag = local(el.tag)
    root.set("xmlns", _SVG_NS)
    return ET.tostring(root, encoding="unicode")


def f2(v: float) -> str:
    s = f"{v:.2f}".rstrip("0").rstrip(".")
    return s if s not in ("", "-0") else "0"


def css_color(c) -> str:
    """'#RRGGBB' / {'color':'#RRGGBB','alpha':0.5} / None → CSS色"""
    if c is None:
        return "transparent"
    if isinstance(c, str):
        return c
    hx, a = c.get("color", "#000000"), c.get("alpha", 1.0)
    if a >= 0.999:
        return hx
    r, g, b = (int(hx[i:i + 2], 16) for i in (1, 3, 5))
    return f"rgba({r},{g},{b},{a:.3f})"


# ---------------------------------------------------------------- geometry
def _poly(pts, close=True):
    d = "M" + " L".join(f"{f2(x)},{f2(y)}" for x, y in pts)
    return d + (" Z" if close else "")


def cmds_to_d(cmds) -> str:
    out = []
    for c in cmds:
        op = c[0]
        if op == "Z":
            out.append("Z")
        else:
            out.append(op + " ".join(f"{f2(c[i])},{f2(c[i + 1])}" for i in range(1, len(c), 2)))
    return " ".join(out)


def geometry(el):
    """戻り値: (paths[(d, filled, stroked)], 線の点列 or None)"""
    g = el.get("geom", "rect")
    w, h = el["w"], el["h"]
    adj = el.get("adj", {})
    ss = min(w, h)

    def a(name, default):
        return adj.get(name, default)

    rect = [(_poly([(0, 0), (w, 0), (w, h), (0, h)]), True, True)]
    if g == "rect":
        return rect, None
    if g == "custom":
        return [(cmds_to_d(p["cmds"]), p.get("fill", True), p.get("stroke", True)) for p in el.get("path", [])], None
    if g == "roundRect":
        r = ss * min(max(a("adj", 16667), 0), 50000) / 100000
        d = (f"M{f2(r)},0 L{f2(w - r)},0 A{f2(r)},{f2(r)} 0 0 1 {f2(w)},{f2(r)} L{f2(w)},{f2(h - r)} "
             f"A{f2(r)},{f2(r)} 0 0 1 {f2(w - r)},{f2(h)} L{f2(r)},{f2(h)} A{f2(r)},{f2(r)} 0 0 1 0,{f2(h - r)} "
             f"L0,{f2(r)} A{f2(r)},{f2(r)} 0 0 1 {f2(r)},0 Z")
        return [(d, True, True)], None
    if g == "ellipse":
        d = (f"M0,{f2(h / 2)} A{f2(w / 2)},{f2(h / 2)} 0 1 1 {f2(w)},{f2(h / 2)} "
             f"A{f2(w / 2)},{f2(h / 2)} 0 1 1 0,{f2(h / 2)} Z")
        return [(d, True, True)], None
    if g == "triangle":
        return [(_poly([(0, h), (w * a("adj", 50000) / 100000, 0), (w, h)]), True, True)], None
    if g == "homePlate":
        v = min(max(a("adj", 50000), 0), 100000 * w / ss if ss else 0)
        x1 = w - ss * v / 100000
        return [(_poly([(0, 0), (x1, 0), (w, h / 2), (x1, h), (0, h)]), True, True)], None
    if g == "chevron":
        v = min(max(a("adj", 50000), 0), 100000 * w / ss if ss else 0)
        x1 = ss * v / 100000
        return [(_poly([(0, 0), (w - x1, 0), (w, h / 2), (w - x1, h), (0, h), (x1, h / 2)]), True, True)], None
    if g in ("flowChartDecision", "diamond"):
        return [(_poly([(0, h / 2), (w / 2, 0), (w, h / 2), (w / 2, h)]), True, True)], None
    if g == "flowChartMerge":
        return [(_poly([(0, 0), (w, 0), (w / 2, h)]), True, True)], None
    if g == "flowChartDocument":
        sx, sy = w / 21600, h / 21600
        d = (f"M0,0 L{f2(w)},0 L{f2(w)},{f2(17322 * sy)} "
             f"C{f2(10800 * sx)},{f2(17322 * sy)} {f2(10800 * sx)},{f2(23922 * sy)} 0,{f2(20172 * sy)} Z")
        return [(d, True, True)], None
    if g in ("can", "flowChartMagneticDisk"):
        ry = ss * a("adj", 25000) / 200000 if g == "can" else h / 6
        rx = w / 2
        body = (f"M0,{f2(ry)} L0,{f2(h - ry)} A{f2(rx)},{f2(ry)} 0 0 0 {f2(w)},{f2(h - ry)} "
                f"L{f2(w)},{f2(ry)} A{f2(rx)},{f2(ry)} 0 0 0 0,{f2(ry)} Z")
        top = f"M0,{f2(ry)} A{f2(rx)},{f2(ry)} 0 0 0 {f2(w)},{f2(ry)}"
        return [(body, True, True), (top, False, True)], None
    if g == "borderCallout1":
        pts = [(w * a("adj2", -8333) / 1e5, h * a("adj1", 18750) / 1e5),
               (w * a("adj4", -38333) / 1e5, h * a("adj3", 112500) / 1e5)]
        return rect + [(_poly(pts, False), False, True)], None
    if g == "borderCallout2":
        pts = [(w * a("adj2", -8333) / 1e5, h * a("adj1", 18750) / 1e5),
               (w * a("adj4", -16667) / 1e5, h * a("adj3", 18750) / 1e5),
               (w * a("adj6", -46667) / 1e5, h * a("adj5", 112500) / 1e5)]
        return rect + [(_poly(pts, False), False, True)], None
    if g in ("line", "straightConnector1"):
        return [], [(0, 0), (w, h)]
    if g == "bentConnector2":
        return [], [(0, 0), (w, 0), (w, h)]
    if g == "bentConnector3":
        x1 = w * a("adj1", 50000) / 1e5
        return [], [(0, 0), (x1, 0), (x1, h), (w, h)]
    if g == "bentConnector4":
        x1, y2 = w * a("adj1", 50000) / 1e5, h * a("adj2", 50000) / 1e5
        return [], [(0, 0), (x1, 0), (x1, y2), (w, y2), (w, h)]
    if g == "bentConnector5":
        x1, y2, x3 = w * a("adj1", 50000) / 1e5, h * a("adj2", 50000) / 1e5, w * a("adj3", 50000) / 1e5
        return [], [(0, 0), (x1, 0), (x1, y2), (x3, y2), (x3, h), (w, h)]
    return rect, None


def text_rect(el):
    """図形ごとの文字領域（presetShapeDefinitions の <a:rect>）。戻り値 (l, t, r, b)"""
    g = el.get("geom", "rect")
    w, h = el["w"], el["h"]
    ss = min(w, h)
    adj = el.get("adj", {})
    if g == "ellipse":
        dx, dy = w / 2 * (1 - math.cos(math.radians(45))), h / 2 * (1 - math.sin(math.radians(45)))
        return dx, dy, w - dx, h - dy
    if g in ("flowChartDecision", "diamond"):
        return w / 4, h / 4, w * 3 / 4, h * 3 / 4
    if g == "flowChartMerge":
        return w / 4, 0, w * 3 / 4, h / 2
    if g == "flowChartDocument":
        return 0, 0, w, h * 17322 / 21600
    if g == "homePlate":
        v = min(max(adj.get("adj", 50000), 0), 100000 * w / ss if ss else 0)
        x1 = w - ss * v / 100000
        return 0, 0, (x1 + w) / 2, h
    if g == "chevron":
        v = min(max(adj.get("adj", 50000), 0), 100000 * w / ss if ss else 0)
        x1 = ss * v / 100000
        return (x1, 0, w - x1, h) if w - x1 > x1 else (0, 0, w, h)
    if g == "triangle":
        x1 = w * adj.get("adj", 50000) / 200000
        return x1, h / 2, x1 + w / 2, h
    if g == "roundRect":
        r = ss * min(max(adj.get("adj", 16667), 0), 50000) / 100000
        il = r * 0.29289
        return il, il, w - il, h - il
    if g == "can":
        y1 = ss * adj.get("adj", 25000) / 200000
        return 0, y1 * 2, w, h - y1
    if g == "flowChartMagneticDisk":
        return 0, h / 3, w, h * 5 / 6
    return 0, 0, w, h


# ---------------------------------------------------------------- html: shape
class _Uid:
    n = 0

    @classmethod
    def next(cls, prefix="c"):
        cls.n += 1
        return f"{prefix}{cls.n}"


def _fill_svg(fill, uid):
    if fill is None:
        return "none", ""
    if isinstance(fill, str) or "color" in fill:
        return css_color(fill), ""
    if "gradient" in fill:
        gr = fill["gradient"]
        rad = math.radians(gr.get("angle", 90))
        x1, y1 = 0.5 - 0.5 * math.cos(rad), 0.5 - 0.5 * math.sin(rad)
        x2, y2 = 0.5 + 0.5 * math.cos(rad), 0.5 + 0.5 * math.sin(rad)
        st = "".join(f'<stop offset="{f2(p)}%" stop-color="{css_color(c)}"/>' for p, c in gr["stops"])
        return (f"url(#{uid})",
                f'<linearGradient id="{uid}" x1="{f2(x1)}" y1="{f2(y1)}" x2="{f2(x2)}" y2="{f2(y2)}">{st}</linearGradient>')
    if "pattern" in fill:
        prst, fg, bg = fill["pattern"], css_color(fill.get("fg", "#000000")), css_color(fill.get("bg", "#FFFFFF"))
        if "DnDiag" in prst or prst == "dnDiag":
            sw = {"ltDnDiag": 1, "dnDiag": 1.5, "dkDnDiag": 2, "wdDnDiag": 3}.get(prst, 1)
            body = f'<path d="M-1,-1 L5,5 M3,-1 L5,1 M-1,3 L1,5" stroke="{fg}" stroke-width="{sw}"/>'
        elif "UpDiag" in prst or prst == "upDiag":
            sw = {"ltUpDiag": 1, "upDiag": 1.5, "dkUpDiag": 2, "wdUpDiag": 3}.get(prst, 1)
            body = f'<path d="M-1,5 L5,-1 M-1,1 L1,-1 M3,5 L5,3" stroke="{fg}" stroke-width="{sw}"/>'
        else:
            body = f'<rect width="1" height="1" fill="{fg}"/>'
        return (f"url(#{uid})", f'<pattern id="{uid}" width="4" height="4" patternUnits="userSpaceOnUse">'
                                 f'<rect width="4" height="4" fill="{bg}"/>{body}</pattern>')
    return "none", ""


def _arrow(end, pts, line, at_start):
    if len(pts) < 2:
        return "", 0
    seq = list(reversed(pts)) if at_start else pts
    tip = seq[-1]
    prev = next((q for q in reversed(seq[:-1]) if abs(q[0] - tip[0]) > 0.01 or abs(q[1] - tip[1]) > 0.01), None)
    if prev is None:
        return "", 0
    ang = math.atan2(tip[1] - prev[1], tip[0] - prev[0])
    lw = line.get("width", 0.75)
    base = max(lw, 2.4)  # 細線でも PowerPoint の矢じりは一定以上の大きさで描かれる
    aw, al = base * ARROW_SZ.get(end.get("w", "med"), 3), base * ARROW_SZ.get(end.get("len", "med"), 3)
    c, s = math.cos(ang), math.sin(ang)

    def tr(dx, dy):
        return (tip[0] + dx * c - dy * s, tip[1] + dx * s + dy * c)
    col = css_color(line.get("color"))
    typ = end.get("type", "triangle")
    if typ in ("triangle", "stealth"):
        p = [tr(0, 0), tr(-al, aw / 2), tr(-al * (0.6 if typ == "stealth" else 1), 0), tr(-al, -aw / 2)]
        return f'<path d="{_poly(p)}" fill="{col}"/>', al * 0.9
    if typ == "arrow":
        p = [tr(-al, aw / 2), tr(0, 0), tr(-al, -aw / 2)]
        return f'<path d="{_poly(p, False)}" fill="none" stroke="{col}" stroke-width="{f2(max(lw, 0.75))}"/>', 0
    if typ == "oval":
        cx, cy = tr(-al / 2, 0)
        return f'<ellipse cx="{f2(cx)}" cy="{f2(cy)}" rx="{f2(al / 2)}" ry="{f2(aw / 2)}" fill="{col}"/>', al / 2
    if typ == "diamond":
        p = [tr(0, 0), tr(-al / 2, aw / 2), tr(-al, 0), tr(-al / 2, -aw / 2)]
        return f'<path d="{_poly(p)}" fill="{col}"/>', al / 2
    return "", 0


def _shorten(pts, d, at_start):
    if d <= 0 or len(pts) < 2:
        return pts
    pts = list(pts)
    i, j = (0, 1) if at_start else (-1, -2)
    a, b = pts[i], pts[j]
    L = math.hypot(b[0] - a[0], b[1] - a[1])
    if L < 0.01:
        return pts
    k = min(d, L * 0.9) / L
    pts[i] = (a[0] + (b[0] - a[0]) * k, a[1] + (b[1] - a[1]) * k)
    return pts


def _shape_svg(el):
    w, h = el["w"], el["h"]
    paths, line_pts = geometry(el)
    fill = el.get("fill") if line_pts is None else None
    line = el.get("line")
    if fill is None and not line:
        return ""
    uid = _Uid.next("p")
    fill_val, defs = _fill_svg(fill, uid)
    stroke = css_color(line.get("color")) if line else "none"
    sw = max(line.get("width", 0.75), 0.75) if line else 0
    dash = ""
    if line and line.get("dash") in DASH:
        dash = f' stroke-dasharray="{" ".join(f2(v * max(sw, 1)) for v in DASH[line["dash"]])}"'
    cap = {"rnd": "round", "sq": "square"}.get((line or {}).get("cap", ""), "butt")
    pad = 12 + sw * 6
    body, extra = [], []
    if line_pts is not None:
        if not line:
            return ""
        pts = line_pts
        if line.get("tail"):
            a, d = _arrow(line["tail"], pts, line, False); extra.append(a); pts = _shorten(pts, d, False)
        if line.get("head"):
            a, d = _arrow(line["head"], pts, line, True); extra.append(a); pts = _shorten(pts, d, True)
        body.append(f'<path d="{_poly(pts, False)}" fill="none" stroke="{stroke}" stroke-width="{f2(sw)}"{dash} '
                    f'stroke-linecap="{cap}"/>')
    else:
        for d, filled, stroked in paths:
            body.append(f'<path d="{d}" fill="{fill_val if filled else "none"}" stroke="{stroke if stroked else "none"}" '
                        f'stroke-width="{f2(sw)}"{dash} stroke-linecap="{cap}"/>')
    fh, fv = el.get("flipH"), el.get("flipV")
    g_open = (f'<g transform="translate({f2(w if fh else 0)},{f2(h if fv else 0)}) scale({-1 if fh else 1},{-1 if fv else 1})">'
              if fh or fv else "<g>")
    filt = ""
    sh = el.get("shadow")
    if sh:
        filt = f' style="left:{f2(-pad)}px;top:{f2(-pad)}px;filter:drop-shadow({f2(sh["dx"])}px {f2(sh["dy"])}px {f2(sh["blur"])}px {css_color(sh["color"])})"'
    else:
        filt = f' style="left:{f2(-pad)}px;top:{f2(-pad)}px"'
    return (f'<svg class="cv-sh" width="{f2(w + 2 * pad)}" height="{f2(h + 2 * pad)}" '
            f'viewBox="{f2(-pad)} {f2(-pad)} {f2(w + 2 * pad)} {f2(h + 2 * pad)}"{filt} aria-hidden="true">'
            f'{"<defs>" + defs + "</defs>" if defs else ""}{g_open}{"".join(body)}{"".join(extra)}</g></svg>')


# ---------------------------------------------------------------- html: text
def _run_css(r):
    fam = []
    for x in r.get("font", []) + ["Meiryo UI"]:
        if x and x not in fam:
            fam.append(x)
    size = r.get("size", 24)
    css = [f"font-size:{f2(size)}px", "font-family:" + ",".join(f"'{x}'" for x in fam),
           f"color:{css_color(r.get('color', '#000000'))}"]
    if r.get("bold"):
        css.append("font-weight:700")
    if r.get("italic"):
        css.append("font-style:italic")
    deco = [d for d, k in (("underline", "underline"), ("line-through", "strike")) if r.get(k)]
    if deco:
        css.append("text-decoration:" + " ".join(deco))
    if r.get("baseline"):
        css.append(f"vertical-align:{r['baseline']:.0f}%;font-size:{f2(size * 0.66)}px")
    if r.get("highlight"):
        css.append(f"background:{css_color(r['highlight'])}")
    return ";".join(css)


def _split_lines(p):
    """改行（br）で区切った行ごとに文字サイズが違う段落は、行ごとの段落に分ける。
    PowerPoint は行ごとに最大文字サイズで行の高さを決めるが、CSS の p は段落全体の最大サイズで決まるため"""
    lines, cur = [], []
    for r in p.get("runs", []):
        if r.get("br"):
            lines.append((cur, r.get("size")))
            cur = []
        else:
            cur.append(r)
    lines.append((cur, None))
    sizes = {max([r.get("size", 24) for r in ln] or [bs or p.get("size") or 24]) for ln, bs in lines}
    if len(lines) < 2 or len(sizes) < 2:
        return [p]
    out = []
    for k, (runs, br_size) in enumerate(lines):
        q = {key: v for key, v in p.items() if key not in ("runs", "spaceBefore", "spaceAfter", "bullet", "indent", "size")}
        q["runs"] = runs
        if not runs:
            q["size"] = br_size or p.get("size") or 24
        if k == 0:
            for key in ("spaceBefore", "bullet", "indent"):
                if key in p:
                    q[key] = p[key]
        if k == len(lines) - 1 and "spaceAfter" in p:
            q["spaceAfter"] = p["spaceAfter"]
        out.append(q)
    return out


def paragraphs_html(paras):
    out = []
    for p in [q for p0 in paras for q in _split_lines(p0)]:
        runs = p.get("runs", [])
        size = p.get("size") or max([r.get("size", 24) for r in runs if not r.get("br")] or [24])
        if p.get("lineHeightPx"):
            lh_css = f"{f2(p['lineHeightPx'])}px"
        else:
            lh_css = f2(p.get("lineHeight", 1.0) * LINE_HEIGHT)
        # 行の高さは p（最大ランの文字サイズ×行間）だけで決め、span は line-height:0 にする（CANVAS_CSS）。
        # フォントの異なる span（箇条書きの Arial と本文の Meiryo UI 等）が混ざると、ブラウザが行ボックスを押し広げるため
        fam = next((r.get("font") for r in runs if r.get("font")), None) or ["Meiryo UI"]
        css = [f"text-align:{ {'l': 'left', 'ctr': 'center', 'r': 'right', 'just': 'justify', 'dist': 'justify'}.get(p.get('align', 'l'), 'left') }",
               f"font-size:{f2(size)}px", "font-family:" + ",".join(f"'{x}'" for x in fam + ["Meiryo UI"]),
               f"line-height:{lh_css}",
               f"margin:{f2(p.get('spaceBefore', 0))}px 0 {f2(p.get('spaceAfter', 0))}px 0",
               f"padding-left:{f2(p.get('marL', 0))}px", f"text-indent:{f2(p.get('indent', 0))}px"]
        inner = []
        bu = p.get("bullet")
        if bu and any((r.get("text") or "").strip() for r in runs):
            bcss = _run_css({"size": bu.get("size", size), "color": bu.get("color", runs[0].get("color") if runs else "#000000"),
                             "font": bu.get("font", runs[0].get("font", []) if runs else []),
                             "bold": runs[0].get("bold") if runs else False})
            inner.append(f'<span class="cv-bu" style="{bcss};width:{f2(max(-p.get("indent", 0), 0))}px">{esc(bu.get("char", "•"))}</span>')
        for r in runs:
            if r.get("br"):
                inner.append("<br>")
            else:
                inner.append(f'<span style="{_run_css(r)}">{esc(r.get("text", ""))}</span>')
        if not any(not r.get("br") for r in runs):
            inner.append(f'<span style="font-size:{f2(size)}px">&#8203;</span>')
        out.append(f'<p style="{";".join(css)}">{"".join(inner)}</p>')
    return "".join(out)


def text_html(tx, w, h):
    if not tx or not tx.get("paragraphs"):
        return ""
    l_in, t_in, r_in, b_in = tx.get("insets", [9.6, 4.8, 9.6, 4.8])
    anchor = tx.get("anchor", "t")
    body = paragraphs_html(tx["paragraphs"])
    pad = f"padding:{f2(t_in)}px {f2(r_in)}px {f2(b_in)}px {f2(l_in)}px"
    rot = f";transform:rotate({f2(tx['rot'])}deg)" if tx.get("rot") else ""
    nowrap = tx.get("wrap") is False
    if tx.get("vert"):
        jc = {"t": "flex-end", "ctr": "center", "b": "flex-start"}.get(anchor, "flex-end")
        return (f'<div class="cv-tx cv-vt" style="width:{f2(w)}px;height:{f2(h)}px;{pad};justify-content:{jc}{rot}">'
                f'<div class="cv-vi{" cv-nw" if nowrap else ""}">{body}</div></div>')
    jc = {"t": "flex-start", "ctr": "center", "b": "flex-end"}.get(anchor, "flex-start")
    if nowrap:
        # 折り返しなしの文字は、段落の揃えに従って箱の左端/中央/右端を基準に伸びる
        al = (tx["paragraphs"][0].get("align", "l") if tx["paragraphs"] else "l")
        ai = {"ctr": "center", "r": "flex-end"}.get(al, "flex-start")
        return (f'<div class="cv-tx cv-nwbox" style="width:{f2(w)}px;height:{f2(h)}px;{pad};justify-content:{jc};align-items:{ai}{rot}">'
                f'<div class="cv-nw">{body}</div></div>')
    inner = "width:auto;margin:0 auto" if tx.get("anchorCtr") else "width:100%"
    return (f'<div class="cv-tx" style="width:{f2(w)}px;height:{f2(h)}px;{pad};justify-content:{jc}{rot}">'
            f'<div style="{inner}">{body}</div></div>')


# ---------------------------------------------------------------- html: other kinds
def _line_css(ln):
    if not ln:
        return "none"
    dash = ln.get("dash") or ""
    st = "dashed" if "ash" in dash else ("dotted" if "ot" in dash else "solid")
    return f"{f2(max(ln.get('width', 0.75), 0.75))}px {st} {css_color(ln.get('color'))}"


def table_html(el):
    cols = el["cols"]
    out = [f'<table class="cv-tb" style="left:{f2(el["x"])}px;top:{f2(el["y"])}px;width:{f2(sum(cols))}px">',
           "<colgroup>" + "".join(f'<col style="width:{f2(c)}px">' for c in cols) + "</colgroup>"]
    for row in el["rows"]:
        out.append(f'<tr style="height:{f2(row["h"])}px">')
        for c in row["cells"]:
            if c.get("merged"):
                continue
            pad = list(c.get("pad", [9.6, 4.8, 9.6, 4.8]))
            # border-collapse では罫線幅の半分がセル寸法に加算される（PowerPoint は加算しない）ため、余白から差し引く
            bds = c.get("borders", {})
            short = [0.0, 0.0, 0.0, 0.0]  # 余白で吸収しきれない分は中身の負マージンで相殺する
            for k, idx in (("l", 0), ("t", 1), ("r", 2), ("b", 3)):
                if bds.get(k):
                    half = max(bds[k].get("width", 0.75), 0.75) / 2
                    short[idx] = max(half - pad[idx], 0)
                    pad[idx] = max(pad[idx] - half, 0)
            css = [f"background:{css_color(c.get('fill')) if c.get('fill') else 'transparent'}",
                   f"padding:{f2(pad[1])}px {f2(pad[2])}px {f2(pad[3])}px {f2(pad[0])}px",
                   "vertical-align:" + {"t": "top", "ctr": "middle", "b": "bottom"}.get(c.get("anchor", "t"), "top")]
            for side, key in (("left", "l"), ("right", "r"), ("top", "t"), ("bottom", "b")):
                if key in c.get("borders", {}):
                    css.append(f"border-{side}:{_line_css(c['borders'][key])}")
            body = paragraphs_html(c.get("paragraphs", []))
            if c.get("vert"):
                body = f'<div class="cv-vc">{body}</div>'
            if any(short):
                body = (f'<div style="margin:{f2(-short[1])}px {f2(-short[2])}px {f2(-short[3])}px {f2(-short[0])}px">'
                        f"{body}</div>")
            attrs = (f' colspan="{c["span"]}"' if c.get("span", 1) > 1 else "") + \
                    (f' rowspan="{c["rowspan"]}"' if c.get("rowspan", 1) > 1 else "")
            cls = " ".join(k for k, on in (("cv-nowrap", c.get("nowrap")), ("cv-clip", c.get("clip"))) if on)
            cls = f' class="{cls}"' if cls else ""
            out.append(f'<td{attrs}{cls} style="{";".join(css)}">{body}</td>')
        out.append("</tr>")
    out.append("</table>")
    return "".join(out)


_SVG_ID_RE = re.compile(r'\bid="([^"]+)"')


def svg_html(el):
    """ベクターアイコン。id 衝突を避けるため要素ごとに接頭辞を付ける"""
    pref = _Uid.next("i") + "_"
    s = el["svg"]  # sanitize_elements() で許可リスト浄化済み
    if not s:
        return ""
    s = _SVG_ID_RE.sub(lambda m: f'id="{pref}{m.group(1)}"', s)
    s = s.replace("url(#", f"url(#{pref}")
    s = re.sub(r'(xlink:href|href)="#', lambda m: f'{m.group(1)}="#{pref}', s)
    s = re.sub(r"<svg\b", f'<svg width="{f2(el["w"])}" height="{f2(el["h"])}" preserveAspectRatio="none"', s, count=1)
    return s


def backdrop_svg(el):
    g = el["grid"]
    gx, gy = g["cols"], g["rows"]
    w, h = el["w"], el["h"]
    cw, ch = w / gx, h / gy
    fid = _Uid.next("bd")
    rects = "".join(
        f'<rect x="{f2(i * cw - 1)}" y="{f2(j * ch - 1)}" width="{f2(cw + 2)}" height="{f2(ch + 2)}" fill="{g["colors"][j * gx + i]}"/>'
        for j in range(gy) for i in range(gx))
    return (f'<svg width="{f2(w)}" height="{f2(h)}" style="position:absolute;left:0;top:0" aria-hidden="true">'
            f'<defs><filter id="{fid}" x="-5%" y="-5%" width="110%" height="110%"><feGaussianBlur stdDeviation="{f2(cw * 0.6)}"/></filter>'
            f'<clipPath id="{fid}c"><rect width="{f2(w)}" height="{f2(h)}"/></clipPath></defs>'
            f'<g clip-path="url(#{fid}c)"><g filter="url(#{fid})">{rects}</g></g></svg>')


def element_html(el):
    k = el.get("kind", "shape")
    if k == "group":
        cx, cy = el.get("origin", [0, 0])
        inner = "".join(element_html(c) for c in el.get("children", []))
        return (f'<div class="cv-grp" style="transform-origin:{f2(cx)}px {f2(cy)}px;'
                f'transform:rotate({f2(el.get("rot", 0))}deg)">{inner}</div>')
    if k == "table":
        return table_html(el)
    box = f'left:{f2(el["x"])}px;top:{f2(el["y"])}px;width:{f2(el["w"])}px;height:{f2(el["h"])}px'
    if el.get("rot"):
        box += f';transform:rotate({f2(el["rot"])}deg)'
    name = f' data-n="{esc(el["name"])}"' if el.get("name") else ""
    title = f' title="{esc(el["alt"])}"' if el.get("alt") else ""
    if k == "svg":
        return f'<div class="cv-s cv-svg"{name}{title} style="{box}">{svg_html(el)}</div>'
    if k == "backdrop":
        return f'<div class="cv-s"{name}{title} style="{box}">{backdrop_svg(el)}</div>'
    if k == "placeholder":
        return f'<div class="cv-s cv-ph"{name} style="{box}"><span>{esc(el.get("label", ""))}</span></div>'
    txt = ""
    if el.get("text"):
        l, t, r, b = text_rect(el)
        txt = text_html(el["text"], r - l, b - t)
        if txt and (l or t):
            txt = f'<div class="cv-tr" style="left:{f2(l)}px;top:{f2(t)}px">{txt}</div>'
    inner = _shape_svg(el) + txt
    if not inner:
        return ""
    return f'<div class="cv-s"{name} style="{box}">{inner}</div>'


def canvas_html(slide) -> str:
    return "".join(element_html(e) for e in sanitize_elements(slide.get("elements", [])))


CANVAS_CSS = """
.cv-s{position:absolute}
.cv-grp{position:absolute;left:0;top:0;width:1280px;height:720px}
.cv-tr{position:absolute}
.cv-sh{position:absolute;overflow:visible;pointer-events:none}
.cv-tx{position:absolute;left:0;top:0;display:flex;flex-direction:column;overflow:visible}
.cv-tx p,.cv-tb p{margin:0;white-space:pre-wrap;word-break:normal;overflow-wrap:anywhere;line-break:strict}
.cv-tx p>span,.cv-tb p>span{line-height:0}
.cv-tb td.cv-clip{overflow:hidden}
.cv-tb td.cv-nowrap{max-width:0}
.cv-tb td.cv-nowrap p{white-space:pre;overflow-wrap:normal}
.cv-nwbox{align-items:center}
.cv-nw{display:flex;flex-direction:column;width:max-content}
.cv-nw p{white-space:pre}
.cv-vt{flex-direction:row}
.cv-vi{writing-mode:vertical-rl;height:100%}
.cv-bu{display:inline-block;text-indent:0}
.cv-tb{position:absolute;border-collapse:collapse;table-layout:fixed}
.cv-tb td{overflow:visible}
.cv-vc{writing-mode:vertical-rl;margin:0 auto}
.cv-svg svg{position:absolute;left:0;top:0}
.cv-ph{border:1px dashed #9a9a9a;background:repeating-linear-gradient(45deg,#f4f4f4 0 10px,#ececec 10px 20px);display:flex;align-items:center;justify-content:center;text-align:center;color:#777;font-size:12px;padding:6px}
"""


# ---------------------------------------------------------------- validation
def validate_canvas(slide, i, fail, warn):
    els = slide.get("elements")
    if not isinstance(els, list):
        fail(f"{i + 1}枚目（type=canvas）: elements は配列で指定してください")

    def walk(lst, path):
        for j, el in enumerate(lst):
            p = f"{path}[{j}]"
            if not isinstance(el, dict):
                fail(f"{i + 1}枚目（type=canvas）: {p} はオブジェクトである必要があります")
            k = el.get("kind", "shape")
            if k not in KINDS:
                fail(f"{i + 1}枚目（type=canvas）: {p}.kind '{k}' は未知です（{', '.join(sorted(KINDS))}）")
            if k == "group":
                walk(el.get("children", []), p + ".children")
                continue
            for key in ("x", "y"):
                if not isinstance(el.get(key), (int, float)):
                    fail(f"{i + 1}枚目（type=canvas）: {p}.{key} が数値ではありません")
            if k == "shape" and el.get("geom", "rect") not in GEOMS:
                warn(f"{i + 1}枚目（type=canvas）: {p}.geom '{el.get('geom')}' は未対応のため矩形で描きます")
            if k == "svg" and "<script" in str(el.get("svg") or "").lower():
                fail(f"{i + 1}枚目（type=canvas）: {p}.svg に script を含めることはできません")
    walk(els, "elements")
    # 色・列挙値・フォント名・パス命令・SVG を描画と同じ規則で検査する（置き換えが起きる値はエラー）
    issues: list = []
    sanitize_elements(els, "elements", issues)
    for msg in issues[:20]:
        fail(f"{i + 1}枚目（type=canvas）: {msg}")


# ---------------------------------------------------------------- pptx
def _emu(v):
    return int(round(v * EMU_PER_PX))


def _hex(c):
    if c is None:
        return None, 1.0
    if isinstance(c, str):
        return c.lstrip("#").upper(), 1.0
    return c.get("color", "#000000").lstrip("#").upper(), c.get("alpha", 1.0)


def _solid_xml(c):
    hx, a = _hex(c)
    alpha = f'<a:alpha val="{int(a * 100000)}"/>' if a < 0.999 else ""
    return f'<a:solidFill><a:srgbClr val="{hx}">{alpha}</a:srgbClr></a:solidFill>'


def _fill_xml(fill):
    if fill is None:
        return "<a:noFill/>"
    if isinstance(fill, str) or "color" in fill:
        return _solid_xml(fill)
    if "gradient" in fill:
        gr = fill["gradient"]
        gs = "".join(f'<a:gs pos="{int(p * 1000)}"><a:srgbClr val="{_hex(c)[0]}"/></a:gs>' for p, c in gr["stops"])
        return f'<a:gradFill><a:gsLst>{gs}</a:gsLst><a:lin ang="{int(gr.get("angle", 90) * 60000)}" scaled="0"/></a:gradFill>'
    if "pattern" in fill:
        return (f'<a:pattFill prst="{fill["pattern"]}"><a:fgClr><a:srgbClr val="{_hex(fill.get("fg", "#000000"))[0]}"/></a:fgClr>'
                f'<a:bgClr><a:srgbClr val="{_hex(fill.get("bg", "#FFFFFF"))[0]}"/></a:bgClr></a:pattFill>')
    return "<a:noFill/>"


def _ln_xml(line, tag="a:ln"):
    if not line:
        return f"<{tag}><a:noFill/></{tag}>"
    dash = f'<a:prstDash val="{line["dash"]}"/>' if line.get("dash") else ""
    ends = ""
    for k, t in (("head", "a:headEnd"), ("tail", "a:tailEnd")):
        if line.get(k):
            e = line[k]
            ends += f'<{t} type="{e.get("type", "triangle")}" w="{e.get("w", "med")}" len="{e.get("len", "med")}"/>'
    cap = f' cap="{line["cap"]}"' if line.get("cap") else ""
    return f'<{tag} w="{_emu(line.get("width", 0.75))}"{cap}>{_solid_xml(line.get("color", "#000000"))}{dash}{ends}</{tag}>'


def _runs_xml(p):
    out = []
    for r in p.get("runs", []):
        if r.get("br"):
            out.append("<a:br/>")
            continue
        attrs = f' lang="ja-JP" sz="{int(round(r.get("size", 24) * 75))}"'
        if r.get("bold"):
            attrs += ' b="1"'
        if r.get("italic"):
            attrs += ' i="1"'
        if r.get("underline"):
            attrs += ' u="sng"'
        if r.get("strike"):
            attrs += ' strike="sngStrike"'
        if r.get("baseline"):
            attrs += f' baseline="{int(r["baseline"] * 1000)}"'
        fonts = r.get("font", [])
        lat = esc(fonts[0]) if fonts else "Meiryo UI"
        ea = esc(fonts[-1]) if fonts else "Meiryo UI"
        out.append(f'<a:r><a:rPr{attrs}>{_solid_xml(r.get("color", "#000000"))}<a:latin typeface="{lat}"/>'
                   f'<a:ea typeface="{ea}"/></a:rPr><a:t>{esc(r.get("text", ""))}</a:t></a:r>')
    return "".join(out)


def _paras_xml(paras):
    out = []
    for p in paras:
        algn = p.get("align", "l")
        ln = f'<a:lnSpc><a:spcPct val="{int(p.get("lineHeight", 1.0) * 100000)}"/></a:lnSpc>'
        if p.get("lineHeightPx"):
            ln = f'<a:lnSpc><a:spcPts val="{int(p["lineHeightPx"] * 75)}"/></a:lnSpc>'
        sp = (f'<a:spcBef><a:spcPts val="{int(p.get("spaceBefore", 0) * 75)}"/></a:spcBef>'
              f'<a:spcAft><a:spcPts val="{int(p.get("spaceAfter", 0) * 75)}"/></a:spcAft>')
        bu = "<a:buNone/>"
        if p.get("bullet"):
            bu = f'<a:buChar char="{esc(p["bullet"].get("char", "•"))}"/>'
        out.append(f'<a:p><a:pPr marL="{_emu(p.get("marL", 0))}" indent="{_emu(p.get("indent", 0))}" algn="{algn}">'
                   f'{ln}{sp}{bu}</a:pPr>{_runs_xml(p)}</a:p>')
    return "".join(out) or "<a:p/>"


def _txbody_xml(tx, tag="p:txBody"):
    if not tx:
        return f'<{tag}><a:bodyPr/><a:lstStyle/><a:p/></{tag}>'
    l_in, t_in, r_in, b_in = tx.get("insets", [9.6, 4.8, 9.6, 4.8])
    vert = ' vert="eaVert"' if tx.get("vert") else ""
    wrap = ' wrap="none"' if tx.get("wrap") is False else ' wrap="square"'
    actr = ' anchorCtr="1"' if tx.get("anchorCtr") else ""
    return (f'<{tag}><a:bodyPr lIns="{_emu(l_in)}" tIns="{_emu(t_in)}" rIns="{_emu(r_in)}" bIns="{_emu(b_in)}"'
            f' anchor="{tx.get("anchor", "t")}"{vert}{wrap}{actr} rtlCol="0"><a:noAutofit/></a:bodyPr><a:lstStyle/>'
            f'{_paras_xml(tx.get("paragraphs", []))}</{tag}>')


def _geom_xml(el):
    g = el.get("geom", "rect")
    if g == "custom":
        w, h = _emu(el["w"]), _emu(el["h"])
        paths = []
        for p in el.get("path", []):
            cmds = []
            for c in p["cmds"]:
                pts = [(_emu(c[i]), _emu(c[i + 1])) for i in range(1, len(c), 2)]
                tag = {"M": "moveTo", "L": "lnTo", "C": "cubicBezTo", "Q": "quadBezTo"}.get(c[0])
                if c[0] == "Z":
                    cmds.append("<a:close/>")
                elif tag:
                    cmds.append(f"<a:{tag}>" + "".join(f'<a:pt x="{x}" y="{y}"/>' for x, y in pts) + f"</a:{tag}>")
            fl = "" if p.get("fill", True) else ' fill="none"'
            sk = "" if p.get("stroke", True) else ' stroke="0"'
            paths.append(f'<a:path w="{w}" h="{h}"{fl}{sk}>{"".join(cmds)}</a:path>')
        return (f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
                f'<a:pathLst>{"".join(paths)}</a:pathLst></a:custGeom>')
    if g not in GEOMS:
        g = "rect"
    av = "".join(f'<a:gd name="{k}" fmla="val {int(v)}"/>' for k, v in el.get("adj", {}).items())
    return f'<a:prstGeom prst="{g}"><a:avLst>{av}</a:avLst></a:prstGeom>'


def _xfrm_xml(el, tag="a:xfrm"):
    rot = f' rot="{int(el.get("rot", 0) * 60000)}"' if el.get("rot") else ""
    fh = ' flipH="1"' if el.get("flipH") else ""
    fv = ' flipV="1"' if el.get("flipV") else ""
    return (f'<{tag}{rot}{fh}{fv}><a:off x="{_emu(el["x"])}" y="{_emu(el["y"])}"/>'
            f'<a:ext cx="{_emu(el["w"])}" cy="{_emu(el["h"])}"/></{tag}>')


NSDECL = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')


def element_pptx_xml(el, sid):
    """要素 → spTree 直下に置く XML 文字列（svg/backdrop は None を返し、呼び出し側で画像化する）"""
    k = el.get("kind", "shape")
    name = esc(el.get("name") or f"canvas {sid}")
    if k == "shape":
        is_line = el.get("geom") in LINE_GEOMS
        fill = "<a:noFill/>" if is_line else _fill_xml(el.get("fill"))
        if is_line:
            return (f'<p:cxnSp {NSDECL}><p:nvCxnSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
                    f'<p:spPr>{_xfrm_xml(el)}{_geom_xml(el)}{fill}{_ln_xml(el.get("line"))}</p:spPr></p:cxnSp>')
        return (f'<p:sp {NSDECL}><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
                f'<p:spPr>{_xfrm_xml(el)}{_geom_xml(el)}{fill}{_ln_xml(el.get("line"))}</p:spPr>'
                f'{_txbody_xml(el.get("text"))}</p:sp>')
    if k == "placeholder":
        tx = {"anchor": "ctr", "paragraphs": [{"align": "ctr", "runs": [{"text": el.get("label", ""), "size": 12, "color": "#777777"}]}]}
        e2 = dict(el, geom="rect", fill="#F0F0F0", line={"color": "#9A9A9A", "width": 1, "dash": "dash"}, text=tx)
        e2["kind"] = "shape"
        return element_pptx_xml(e2, sid)
    if k == "table":
        cols = "".join(f'<a:gridCol w="{_emu(c)}"/>' for c in el["cols"])
        rows = []
        for row in el["rows"]:
            cells = []
            for c in row["cells"]:
                attrs = ""
                if c.get("span", 1) > 1:
                    attrs += f' gridSpan="{c["span"]}"'
                if c.get("rowspan", 1) > 1:
                    attrs += f' rowSpan="{c["rowspan"]}"'
                if c.get("merged") == "h":
                    attrs += ' hMerge="1"'
                elif c.get("merged") == "v":
                    attrs += ' vMerge="1"'
                pad = c.get("pad", [9.6, 4.8, 9.6, 4.8])
                bd = "".join(_ln_xml(c.get("borders", {}).get(s), t) for s, t in (("l", "a:lnL"), ("r", "a:lnR"), ("t", "a:lnT"), ("b", "a:lnB")))
                tx = {"paragraphs": c.get("paragraphs", [])}
                body = _txbody_xml(tx, "a:txBody").replace("<a:noAutofit/>", "")
                cells.append(f'<a:tc{attrs}>{body}<a:tcPr marL="{_emu(pad[0])}" marR="{_emu(pad[2])}" marT="{_emu(pad[1])}" '
                             f'marB="{_emu(pad[3])}" anchor="{c.get("anchor", "t")}">{bd}{_fill_xml(c.get("fill"))}</a:tcPr></a:tc>')
            rows.append(f'<a:tr h="{_emu(row["h"])}">{"".join(cells)}</a:tr>')
        return (f'<p:graphicFrame {NSDECL}><p:nvGraphicFramePr><p:cNvPr id="{sid}" name="{name}"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
                f'{_xfrm_xml(dict(el, w=sum(el["cols"]), h=sum(r["h"] for r in el["rows"])), "p:xfrm")}'
                f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
                f'<a:tbl><a:tblPr/><a:tblGrid>{cols}</a:tblGrid>{"".join(rows)}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>')
    return None


def flatten(elements):
    """group を展開した要素列（PPTX 出力用。group の回転は子に加算する簡易処理）"""
    for el in elements:
        if el.get("kind") == "group":
            for c in flatten(el.get("children", [])):
                if el.get("rot") and c.get("kind") != "table":
                    c = dict(c)
                    cx, cy = el.get("origin", [0, 0])
                    ang = math.radians(el["rot"])
                    mx, my = c["x"] + c.get("w", 0) / 2 - cx, c["y"] + c.get("h", 0) / 2 - cy
                    nx, ny = cx + mx * math.cos(ang) - my * math.sin(ang), cy + mx * math.sin(ang) + my * math.cos(ang)
                    c["x"], c["y"] = nx - c.get("w", 0) / 2, ny - c.get("h", 0) / 2
                    c["rot"] = c.get("rot", 0) + el["rot"]
                yield c
        else:
            yield el


def add_canvas_to_pptx(slide, elements, png_for_svg=None, png_for_backdrop=None):
    """python-pptx の slide に canvas 要素を追加する"""
    from pptx.oxml import parse_xml
    from pptx.util import Emu
    tree = slide.shapes._spTree
    sid = 100
    for el in flatten(sanitize_elements(elements)):
        sid += 1
        k = el.get("kind", "shape")
        if k in ("svg", "backdrop"):
            png = (png_for_svg if k == "svg" else png_for_backdrop)
            path = png(el) if png else None
            if path:
                pic = slide.shapes.add_picture(str(path), Emu(_emu(el["x"])), Emu(_emu(el["y"])),
                                               Emu(_emu(el["w"])), Emu(_emu(el["h"])))
                if el.get("rot"):
                    pic.rotation = el["rot"]
            continue
        xml = element_pptx_xml(el, sid)
        if xml:
            tree.append(parse_xml(xml))
