#!/usr/bin/env python3
"""既存 PPTX を canvas 型スライドの deck.json に取り込む（見た目をほぼそのまま再現する）。

  python tools/pptx_import.py <src.pptx> --into <deck_dir> [--id <deck_id>] [--slides 1 5 28]

- 図形（矩形・矢羽根・フローチャート記号・コネクタ等）は座標・色・線種・矢印をそのまま canvas 要素にする
- 文字はプレースホルダー/マスター/既定テキストスタイルの継承を解決し、ランごとの書式を持たせる
- 表はセル結合・塗り・罫線（表スタイル含む）を解決して table 要素にする
- SVG アイコンはベクターのまま埋め込む。ラスタ画像（スクリーンショット等）は埋め込まず placeholder にする
  （全面背景の写真だけは色格子＋ぼかしの backdrop で近似）
- 埋め込み Excel（OLE）は、ブックの内容を表として取り込む
"""
from __future__ import annotations

import argparse
import colorsys
import io
import json
import math
import re
import sys
import zipfile
from copy import deepcopy
from pathlib import Path

from lxml import etree

# 取り込み元の PPTX は信用しない: 外部実体・ネットワーク参照・巨大ツリーを解決しないパーサで読む
_SAFE_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False, load_dtd=False)


def _xml(data):
    return etree.fromstring(data, _SAFE_PARSER)

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]
R = "{%s}" % NS["r"]
EMU_PX = 9525.0
PT_PX = 96.0 / 72.0
PRST = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF",
        "yellow": "FFFF00", "gray": "808080", "darkGray": "A9A9A9", "lightGray": "D3D3D3"}
COLOR_TAGS = ("srgbClr", "schemeClr", "sysClr", "prstClr", "scrgbClr")
# PowerPoint の既定の表スタイル「中間スタイル 2 - アクセント 1」。PowerPoint は組み込みスタイルを ID だけで適用するため、
# tableStyles.xml に定義が無い PPTX（python-pptx 製など）でも同じ見た目にするための既定定義
_LN_LT1 = '<a:ln w="12700" cmpd="sng"><a:solidFill><a:schemeClr val="lt1"/></a:solidFill></a:ln>'
_BOLD_LT1 = '<a:tcTxStyle b="on"><a:fontRef idx="minor"><a:prstClr val="black"/></a:fontRef><a:schemeClr val="lt1"/></a:tcTxStyle>'
BUILTIN_TABLE_STYLES = {
    "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}": (
        f'<a:tblStyle xmlns:a="{NS["a"]}" styleId="{{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}}">'
        '<a:wholeTbl><a:tcTxStyle><a:fontRef idx="minor"><a:prstClr val="black"/></a:fontRef><a:schemeClr val="dk1"/></a:tcTxStyle>'
        f'<a:tcStyle><a:tcBdr><a:left>{_LN_LT1}</a:left><a:right>{_LN_LT1}</a:right><a:top>{_LN_LT1}</a:top>'
        f'<a:bottom>{_LN_LT1}</a:bottom><a:insideH>{_LN_LT1}</a:insideH><a:insideV>{_LN_LT1}</a:insideV></a:tcBdr>'
        '<a:fill><a:solidFill><a:schemeClr val="accent1"><a:tint val="20000"/></a:schemeClr></a:solidFill></a:fill></a:tcStyle></a:wholeTbl>'
        '<a:band1H><a:tcStyle><a:tcBdr/><a:fill><a:solidFill><a:schemeClr val="accent1"><a:tint val="40000"/></a:schemeClr></a:solidFill></a:fill></a:tcStyle></a:band1H>'
        '<a:band1V><a:tcStyle><a:tcBdr/><a:fill><a:solidFill><a:schemeClr val="accent1"><a:tint val="40000"/></a:schemeClr></a:solidFill></a:fill></a:tcStyle></a:band1V>'
        f'<a:lastCol>{_BOLD_LT1}<a:tcStyle><a:tcBdr/><a:fill><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:fill></a:tcStyle></a:lastCol>'
        f'<a:firstCol>{_BOLD_LT1}<a:tcStyle><a:tcBdr/><a:fill><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:fill></a:tcStyle></a:firstCol>'
        f'<a:lastRow>{_BOLD_LT1}<a:tcStyle><a:tcBdr><a:top><a:ln w="38100" cmpd="sng"><a:solidFill><a:schemeClr val="lt1"/></a:solidFill></a:ln></a:top></a:tcBdr>'
        '<a:fill><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:fill></a:tcStyle></a:lastRow>'
        f'<a:firstRow>{_BOLD_LT1}<a:tcStyle><a:tcBdr><a:bottom><a:ln w="38100" cmpd="sng"><a:solidFill><a:schemeClr val="lt1"/></a:solidFill></a:ln></a:bottom></a:tcBdr>'
        '<a:fill><a:solidFill><a:schemeClr val="accent1"/></a:solidFill></a:fill></a:tcStyle></a:firstRow></a:tblStyle>'),
}
WARNINGS: list = []


def warn(msg):
    WARNINGS.append(msg)
    print(f"warning: {msg}", file=sys.stderr)


def px(v) -> float:
    return round(int(v) / EMU_PX, 2)


def r2(v: float) -> float:
    return round(v, 2)


def lname(el) -> str:
    return etree.QName(el).localname


def color_child(el):
    if el is None:
        return None
    for ch in el:
        if lname(ch) in COLOR_TAGS:
            return ch
    return None


def first_attr(els, name, default=None):
    for el in els:
        if el is not None and el.get(name) is not None:
            return el.get(name)
    return default


def first_child(els, tag):
    for el in els:
        if el is None:
            continue
        ch = el.find(A + tag)
        if ch is not None:
            return ch
    return None


def _lin(c):
    return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4


def _srgb(c):
    c = min(max(c, 0.0), 1.0)
    return c * 12.92 if c <= 0.0031308 else 1.055 * c ** (1 / 2.4) - 0.055


def ccol(c):
    """(hex, alpha) → canvas の色表現"""
    if c is None:
        return None
    hx, a = c
    return hx if a >= 0.999 else {"color": hx, "alpha": round(a, 3)}


class Theme:
    def __init__(self, theme_xml: bytes, clrmap: dict):
        t = _xml(theme_xml)
        self.colors = {}
        for el in t.find(f"{A}themeElements/{A}clrScheme"):
            ch = el[0]
            self.colors[lname(el)] = ch.get("lastClr") if lname(ch) == "sysClr" else ch.get("val")
        fs = t.find(f"{A}themeElements/{A}fontScheme")
        self.fonts = {
            "+mj-lt": fs.find(f"{A}majorFont/{A}latin").get("typeface"),
            "+mj-ea": fs.find(f"{A}majorFont/{A}ea").get("typeface"),
            "+mn-lt": fs.find(f"{A}minorFont/{A}latin").get("typeface"),
            "+mn-ea": fs.find(f"{A}minorFont/{A}ea").get("typeface"),
        }
        fmt = t.find(f"{A}themeElements/{A}fmtScheme")
        self.fill_styles = list(fmt.find(f"{A}fillStyleLst"))
        self.line_styles = list(fmt.find(f"{A}lnStyleLst"))
        self.bg_styles = list(fmt.find(f"{A}bgFillStyleLst"))
        self.clrmap = clrmap

    def font(self, face):
        if not face:
            return None
        face = face.replace("+mn-cs", "+mn-lt").replace("+mj-cs", "+mj-lt")
        return self.fonts.get(face, face)

    def color(self, el, ph=None):
        """色要素を (#RRGGBB, alpha) にする。ph は phClr の置換色"""
        if el is None:
            return None
        tag = lname(el)
        alpha = 1.0
        if tag == "srgbClr":
            hx = el.get("val")
        elif tag == "schemeClr":
            v = el.get("val")
            if v == "phClr":
                if ph is None:
                    return None
                hx, alpha = ph[0].lstrip("#"), ph[1]
            else:
                hx = self.colors.get(self.clrmap.get(v, v), "000000")
        elif tag == "sysClr":
            hx = el.get("lastClr") or "000000"
        elif tag == "prstClr":
            hx = PRST.get(el.get("val"), "000000")
        elif tag == "scrgbClr":
            rr, gg, bb = (int(el.get(k)) / 100000 for k in ("r", "g", "b"))
            hx = "%02X%02X%02X" % tuple(round(_srgb(c) * 255) for c in (rr, gg, bb))
        else:
            return None
        r, g, b = (int(hx[i:i + 2], 16) / 255 for i in (0, 2, 4))
        for m in el:
            t, v = lname(m), int(m.get("val", "0")) / 100000
            if t == "alpha":
                alpha *= v
            elif t in ("lumMod", "lumOff", "satMod"):
                h, l, s = colorsys.rgb_to_hls(r, g, b)
                if t == "lumMod":
                    l *= v
                elif t == "lumOff":
                    l += v
                else:
                    s *= v
                r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0), 1), min(max(s, 0), 1))
            elif t == "tint":
                r, g, b = (_srgb(_lin(c) * v + (1 - v)) for c in (r, g, b))
            elif t == "shade":
                r, g, b = (_srgb(_lin(c) * v) for c in (r, g, b))
        return "#%02X%02X%02X" % tuple(round(min(max(c, 0), 1) * 255) for c in (r, g, b)), alpha


# ---------------------------------------------------------------- line / fill
class Line:
    def __init__(self):
        self.color = None
        self.width = 0.0
        self.dash = None
        self.head = None
        self.tail = None
        self.cap = None
        self.nofill = False

    def merge(self, ln, theme, ph=None):
        if ln is None:
            return
        if ln.get("w") is not None:
            self.width = int(ln.get("w")) / EMU_PX
        if ln.get("cap"):
            self.cap = ln.get("cap")
        for ch in ln:
            t = lname(ch)
            if t == "noFill":
                self.nofill = True
            elif t == "solidFill":
                self.color = theme.color(color_child(ch), ph); self.nofill = False
            elif t == "gradFill":
                gs = ch.find(f"{A}gsLst/{A}gs")
                if gs is not None:
                    self.color = theme.color(color_child(gs), ph); self.nofill = False
            elif t == "prstDash":
                self.dash = None if ch.get("val") == "solid" else ch.get("val")
            elif t in ("headEnd", "tailEnd"):
                end = None
                if ch.get("type", "none") != "none":
                    end = {"type": ch.get("type")}
                    for k in ("w", "len"):
                        if ch.get(k) and ch.get(k) != "med":
                            end[k] = ch.get(k)
                setattr(self, "head" if t == "headEnd" else "tail", end)

    def to_json(self):
        if self.nofill or self.color is None:
            return None
        out = {"color": ccol(self.color), "width": r2(max(self.width, 0.75) if self.width else 0.75)}
        if self.dash:
            out["dash"] = self.dash
        if self.cap and self.cap != "flat":
            out["cap"] = self.cap
        if self.head:
            out["head"] = self.head
        if self.tail:
            out["tail"] = self.tail
        return out


# ---------------------------------------------------------------- importer
class Importer:
    def __init__(self, src: Path):
        from pptx import Presentation
        self.src = src
        self.prs = Presentation(str(src))
        self.zip = zipfile.ZipFile(src)
        master = self.prs.slide_masters[0]
        self.base_map = dict(master._element.find(f"{P}clrMap").attrib)
        theme_part = next(r.target_part for r in master.part.rels.values() if r.reltype.endswith("/theme"))
        self.theme = Theme(theme_part.blob, dict(self.base_map))
        self.W = int(self.prs.slide_width) / EMU_PX
        self.H = int(self.prs.slide_height) / EMU_PX
        pres = _xml(self.zip.read("ppt/presentation.xml"))
        # 既定テキストスタイルの色は「最も弱い既定値」なので、図形スタイル(fontRef)や表スタイルの文字色に
        # 負けるよう、色指定を取り除いた版を継承チェーンに入れる（色の最終既定はテーマの tx1）
        self.default_text = deepcopy(pres.find(f"{P}defaultTextStyle"))
        if self.default_text is not None:
            for sf in list(self.default_text.iter(f"{A}solidFill")):
                sf.getparent().remove(sf)
        self.tbl_styles = {k: _xml(v) for k, v in BUILTIN_TABLE_STYLES.items()}
        try:
            for st in _xml(self.zip.read("ppt/tableStyles.xml")).findall(f"{A}tblStyle"):
                self.tbl_styles[st.get("styleId")] = st
        except KeyError:
            pass
        self._grp_fill = None

    # ---- placeholder inheritance
    @staticmethod
    def ph_info(sp):
        ph = sp.find(f".//{P}nvPr/{P}ph")
        if ph is None:
            return None
        return ph.get("type", "body"), ph.get("idx")

    @classmethod
    def find_ph(cls, tree, typ, idx):
        cands = [(sp, cls.ph_info(sp)) for sp in tree.iter(f"{P}sp") if cls.ph_info(sp) is not None]
        if idx is not None:
            for sp, (t, i) in cands:
                if i == idx:
                    return sp
        norm = {"ctrTitle": "title", "subTitle": "body", "obj": "body"}
        for sp, (t, i) in cands:
            if norm.get(t, t) == norm.get(typ, typ):
                return sp
        return None

    def ph_chain(self, sp, ctx):
        info = self.ph_info(sp)
        if info is None:
            return []
        typ, idx = info
        lay = self.find_ph(ctx["layout"], typ, idx)
        lt = self.ph_info(lay)[0] if lay is not None else typ
        mas = self.find_ph(ctx["master"], lt, None)
        return [x for x in (lay, mas) if x is not None]

    def master_text_style(self, typ, master_el):
        tx = master_el.find(f"{P}txStyles")
        if typ in ("title", "ctrTitle"):
            return tx.find(f"{P}titleStyle")
        if typ in ("body", "obj", "subTitle"):
            return tx.find(f"{P}bodyStyle")
        return tx.find(f"{P}otherStyle")

    # ---- fill / line
    def grad(self, g, ph=None):
        stops = [[round(int(gs.get("pos")) / 1000, 1), ccol(self.theme.color(color_child(gs), ph))]
                 for gs in g.findall(f"{A}gsLst/{A}gs")]
        lin = g.find(f"{A}lin")
        ang = int(lin.get("ang", "0")) / 60000 if lin is not None else 90
        return {"gradient": {"angle": round(ang, 1), "stops": sorted(stops, key=lambda s: s[0])}}

    def fill_of(self, sppr, style, ph_srcs=()):
        for src in [sppr] + [s.find(f"{P}spPr") for s in ph_srcs]:
            if src is None:
                continue
            for ch in src:
                t = lname(ch)
                if t == "noFill":
                    return None
                if t == "solidFill":
                    return ccol(self.theme.color(color_child(ch)))
                if t == "gradFill":
                    return self.grad(ch)
                if t == "pattFill":
                    return {"pattern": ch.get("prst"),
                            "fg": ccol(self.theme.color(color_child(ch.find(f"{A}fgClr")))) or "#000000",
                            "bg": ccol(self.theme.color(color_child(ch.find(f"{A}bgClr")))) or "#FFFFFF"}
                if t == "grpFill":
                    return self._grp_fill
        if style is not None:
            fr = style.find(f"{A}fillRef")
            if fr is not None:
                idx = int(fr.get("idx", "0"))
                phc = self.theme.color(color_child(fr))
                if idx == 0 or phc is None:
                    return None
                lst = self.theme.fill_styles if idx < 1000 else self.theme.bg_styles
                k = (idx if idx < 1000 else idx - 1000) - 1
                if 0 <= k < len(lst):
                    st = lst[k]
                    if lname(st) == "solidFill":
                        return ccol(self.theme.color(color_child(st), phc))
                    if lname(st) == "gradFill":
                        return self.grad(st, phc)
                    if lname(st) == "noFill":
                        return None
                return ccol(phc)
        return None

    def line_of(self, sppr, style, ph_srcs=()):
        ls = Line()
        if style is not None:
            lr = style.find(f"{A}lnRef")
            if lr is not None and int(lr.get("idx", "0")) > 0:
                k = int(lr.get("idx")) - 1
                if k < len(self.theme.line_styles):
                    ls.merge(self.theme.line_styles[k], self.theme, self.theme.color(color_child(lr)))
        for src in reversed([sppr] + [s.find(f"{P}spPr") for s in ph_srcs]):
            if src is not None:
                ls.merge(src.find(f"{A}ln"), self.theme)
        return ls.to_json()

    # ---- text
    def runs_of(self, p, rbase, font_color, scale, slide_no):
        theme = self.theme
        out = []

        def fmt(rpr):
            chain = [rpr] + rbase
            sz = first_attr(chain, "sz")
            f = {"size": r2((int(sz) / 100 if sz else 18) * PT_PX * scale)}
            if first_attr(chain, "b") in ("1", "true"):
                f["bold"] = True
            if first_attr(chain, "i") in ("1", "true"):
                f["italic"] = True
            u = first_attr(chain, "u")
            if u and u != "none":
                f["underline"] = True
            s = first_attr(chain, "strike")
            if s and s != "noStrike":
                f["strike"] = True
            b = first_attr(chain, "baseline")
            if b and int(b) != 0:
                f["baseline"] = int(b) / 1000
            color = None
            for el in chain:
                if el is None:
                    continue
                sf = el.find(f"{A}solidFill")
                if sf is not None:
                    color = theme.color(color_child(sf)); break
                gf = el.find(f"{A}gradFill")
                if gf is not None:
                    gs = gf.find(f"{A}gsLst/{A}gs")
                    color = theme.color(color_child(gs)); break
                if el.find(f"{A}noFill") is not None:
                    color = ("#000000", 0.0); break
            f["color"] = ccol(color or font_color or theme.color(etree.Element(A + "schemeClr", val="tx1")))
            fonts = []
            for tag in ("latin", "ea"):
                face = theme.font(first_attr([c.find(A + tag) if c is not None else None for c in chain], "typeface"))
                if face and face not in fonts:
                    fonts.append(face)
            f["font"] = fonts
            for el in chain:
                if el is not None and el.find(f"{A}highlight") is not None:
                    f["highlight"] = ccol(theme.color(color_child(el.find(f"{A}highlight")))); break
            return f

        for r in p:
            t = lname(r)
            if t not in ("r", "br", "fld"):
                continue
            f = fmt(r.find(f"{A}rPr"))
            if t == "br":
                out.append({"br": True, "size": f["size"]})
                continue
            text = r.findtext(f"{A}t") or ""
            if t == "fld" and r.get("type") == "slidenum":
                text = str(slide_no)
            f["text"] = text
            out.append(f)
        if not any("text" in x for x in out):
            return out, fmt(p.find(f"{A}endParaRPr"))["size"]
        return out, max(x["size"] for x in out)

    def text_of(self, txbody, chain, body_props, font_color, slide_no, bold=False):
        if txbody is None:
            return None
        paras = txbody.findall(f"{A}p")
        if not any((p.xpath("string(.)") or "").strip() for p in paras):
            return None
        bp = body_props

        def ins(k, d):
            return round(int(first_attr(bp, k, d)) / EMU_PX, 2)
        tx = {"insets": [ins("lIns", 91440), ins("tIns", 45720), ins("rIns", 91440), ins("bIns", 45720)]}
        anchor = first_attr(bp, "anchor", "t")
        if anchor != "t":
            tx["anchor"] = anchor
        if first_attr(bp, "anchorCtr", "0") in ("1", "true"):
            tx["anchorCtr"] = True
        if first_attr(bp, "wrap", "square") == "none":
            tx["wrap"] = False
        if first_attr(bp, "vert", "horz") in ("eaVert", "vert", "wordArtVert", "mongolianVert"):
            tx["vert"] = True
        if bp and bp[0] is not None and bp[0].get("rot"):
            tx["rot"] = int(bp[0].get("rot")) / 60000
        scale, ln_red = 1.0, 0.0
        for b in bp:
            if b is None:
                continue
            na = b.find(f"{A}normAutofit")
            if na is not None:
                scale = int(na.get("fontScale", "100000")) / 100000
                ln_red = int(na.get("lnSpcReduction", "0")) / 100000
                break
            if b.find(f"{A}noAutofit") is not None or b.find(f"{A}spAutoFit") is not None:
                break
        lsts = [txbody.find(f"{A}lstStyle")] + chain
        out = []
        autonum = {}
        for p in paras:
            ppr = p.find(f"{A}pPr")
            lvl = int(ppr.get("lvl", "0")) if ppr is not None else 0
            pchain = [ppr] + [l.find(f"{A}lvl{lvl + 1}pPr") if l is not None else None for l in lsts]
            rbase = [x.find(f"{A}defRPr") if x is not None else None for x in pchain]
            runs, max_fs = self.runs_of(p, rbase, font_color, scale, slide_no)
            if bold:
                for r in runs:
                    if "text" in r and "bold" not in r:
                        r["bold"] = True
            pj = {"runs": runs}
            if not any("text" in r for r in runs):
                pj["size"] = max_fs
            algn = first_attr(pchain, "algn", "l")
            if algn != "l":
                pj["align"] = algn
            mar_l, indent = px(first_attr(pchain, "marL", 0)), px(first_attr(pchain, "indent", 0))
            if mar_l:
                pj["marL"] = mar_l
            if indent:
                pj["indent"] = indent
            ln = first_child(pchain, "lnSpc")
            lh = 1.0
            if ln is not None:
                pct, pts = ln.find(f"{A}spcPct"), ln.find(f"{A}spcPts")
                if pct is not None:
                    lh = int(pct.get("val")) / 100000
                elif pts is not None:
                    pj["lineHeightPx"] = r2(int(pts.get("val")) / 100 * PT_PX)
            lh = max(lh - ln_red, 0.3)
            if abs(lh - 1.0) > 0.001 and "lineHeightPx" not in pj:
                pj["lineHeight"] = round(lh, 3)
            for key, tag in (("spaceBefore", "spcBef"), ("spaceAfter", "spcAft")):
                el = first_child(pchain, tag)
                v = 0
                if el is not None:
                    if el.find(f"{A}spcPts") is not None:
                        v = int(el.find(f"{A}spcPts").get("val")) / 100 * PT_PX
                    elif el.find(f"{A}spcPct") is not None:
                        v = int(el.find(f"{A}spcPct").get("val")) / 100000 * max_fs * 1.2
                if v:
                    pj[key] = r2(v)
            bu_kind = None
            for el in pchain:
                if el is None:
                    continue
                for t in ("buNone", "buChar", "buAutoNum"):
                    if el.find(A + t) is not None:
                        bu_kind = t; break
                if bu_kind:
                    break
            if bu_kind in ("buChar", "buAutoNum") and (p.xpath("string(.)") or "").strip():
                if bu_kind == "buChar":
                    ch = first_child(pchain, "buChar").get("char") or "•"
                else:
                    an = first_child(pchain, "buAutoNum")
                    key = (lvl, an.get("type"))
                    autonum[key] = autonum.get(key, int(an.get("startAt", "1")) - 1) + 1
                    n = autonum[key]
                    ch = {"arabicPeriod": f"{n}.", "arabicParenR": f"{n})", "arabicParenBoth": f"({n})",
                          "circleNumDbPlain": chr(0x2460 + n - 1)}.get(key[1] or "arabicPeriod", f"{n}.")
                bu = {"char": ch}
                first_run = next((r for r in runs if "text" in r), {})
                bclr = first_child(pchain, "buClr")
                if bclr is not None:
                    bu["color"] = ccol(self.theme.color(color_child(bclr)))
                bsz = first_child(pchain, "buSzPct")
                if bsz is not None:
                    bu["size"] = r2(first_run.get("size", max_fs) * int(bsz.get("val")) / 100000)
                bfont = first_child(pchain, "buFont")
                if bfont is not None and bfont.get("typeface") and not bfont.get("typeface").startswith("+"):
                    bu["font"] = [bfont.get("typeface")]
                pj["bullet"] = bu
            out.append(pj)
        tx["paragraphs"] = out
        return tx

    # ---- tree walk
    @staticmethod
    def map_box(xf, x, y, w, h):
        ox, oy, sx, sy, cx, cy = xf
        return ox + (x - cx) * sx, oy + (y - cy) * sy, w * sx, h * sy

    def box_of(self, x, xf):
        off, ext = x.find(f"{A}off"), x.find(f"{A}ext")
        X, Y, W, H = self.map_box(xf, int(off.get("x")) / EMU_PX, int(off.get("y")) / EMU_PX,
                                  int(ext.get("cx")) / EMU_PX, int(ext.get("cy")) / EMU_PX)
        return r2(X), r2(Y), r2(W), r2(H)

    def walk(self, tree, ctx, xf):
        out = []
        for el in tree:
            t = lname(el)
            try:
                if t in ("sp", "cxnSp"):
                    e = self.shape(el, ctx, xf)
                    if e:
                        out.append(e)
                elif t == "grpSp":
                    out.extend(self.group(el, ctx, xf))
                elif t == "pic":
                    e = self.picture(el, ctx, xf)
                    if e:
                        out.append(e)
                elif t == "graphicFrame":
                    e = self.frame(el, ctx, xf)
                    if e:
                        out.append(e)
                elif t == "AlternateContent":
                    ch = el.find(f"{{{NS['mc']}}}Choice")
                    if ch is None:
                        ch = el.find(f"{{{NS['mc']}}}Fallback")
                    if ch is not None:
                        out.extend(self.walk(ch, ctx, xf))
            except Exception as ex:  # noqa: BLE001
                nm = el.find(f".//{P}cNvPr")
                warn(f"スライド{ctx['no']}: {nm.get('name') if nm is not None else t} の取り込みに失敗: {ex!r}")
        return out

    def group(self, g, ctx, xf):
        gx = g.find(f"{P}grpSpPr/{A}xfrm")
        if gx is None:
            return self.walk(g, ctx, xf)
        ox, oy, ow, oh = self.box_of(gx, xf)
        ch_off, ch_ext = gx.find(f"{A}chOff"), gx.find(f"{A}chExt")
        cw, chh = int(ch_ext.get("cx")) / EMU_PX, int(ch_ext.get("cy")) / EMU_PX
        nxf = (ox, oy, ow / cw if cw else 1, oh / chh if chh else 1,
               int(ch_off.get("x")) / EMU_PX, int(ch_off.get("y")) / EMU_PX)
        saved = self._grp_fill
        self._grp_fill = self.fill_of(g.find(f"{P}grpSpPr"), None)
        children = self.walk(g, ctx, nxf)
        self._grp_fill = saved
        rot = int(gx.get("rot", "0")) / 60000
        if rot:
            return [{"kind": "group", "rot": round(rot, 2), "origin": [r2(ox + ow / 2), r2(oy + oh / 2)], "children": children}]
        return children

    def shape(self, sp, ctx, xf):
        info = self.ph_info(sp)
        if info and ctx.get("skip_ph"):
            return None
        ph_srcs = self.ph_chain(sp, ctx) if info else []
        x = None
        for src in [sp] + ph_srcs:
            x = src.find(f"{P}spPr/{A}xfrm")
            if x is not None and x.find(f"{A}off") is not None:
                break
            x = None
        if x is None:
            return None
        X, Y, W, H = self.box_of(x, xf)
        el = {"x": X, "y": Y, "w": W, "h": H}
        name = sp.find(f".//{P}cNvPr")
        if name is not None and name.get("name"):
            el["name"] = name.get("name")
        rot = int(x.get("rot", "0")) / 60000
        if rot:
            el["rot"] = round(rot, 2)
        if x.get("flipH") == "1":
            el["flipH"] = True
        if x.get("flipV") == "1":
            el["flipV"] = True
        sppr = sp.find(f"{P}spPr")
        style = sp.find(f"{P}style")
        geom = sppr.find(f"{A}prstGeom") if sppr is not None else None
        cust = sppr.find(f"{A}custGeom") if sppr is not None else None
        if geom is None and cust is None:
            for s in ph_srcs:
                geom = s.find(f"{P}spPr/{A}prstGeom")
                if geom is not None:
                    break
        if cust is not None:
            el["geom"] = "custom"
            el["path"] = self.cust_path(cust, W, H)
        else:
            prst = geom.get("prst") if geom is not None else "rect"
            if prst != "rect":
                el["geom"] = prst
            from canvas_engine import GEOMS  # type: ignore
            if prst not in GEOMS:
                warn(f"スライド{ctx['no']}: 未対応の図形 {prst}（{el.get('name', '')}）は矩形で描きます")
            av = geom.find(f"{A}avLst") if geom is not None else None
            if av is not None:
                adj = {}
                for gd in av:
                    m = re.match(r"val\s+(-?\d+)", gd.get("fmla", ""))
                    if m:
                        adj[gd.get("name")] = int(m.group(1))
                if adj:
                    el["adj"] = adj
        from canvas_engine import LINE_GEOMS  # type: ignore
        if el.get("geom") not in LINE_GEOMS:
            fill = self.fill_of(sppr, style, ph_srcs)
            if fill is not None:
                el["fill"] = fill
        line = self.line_of(sppr, style, ph_srcs)
        if line:
            el["line"] = line
        eff = sppr.find(f"{A}effectLst/{A}outerShdw") if sppr is not None else None
        if eff is not None:
            dist = int(eff.get("dist", "0")) / EMU_PX
            dr = math.radians(int(eff.get("dir", "0")) / 60000)
            el["shadow"] = {"dx": r2(dist * math.cos(dr)), "dy": r2(dist * math.sin(dr)),
                            "blur": r2(int(eff.get("blurRad", "0")) / EMU_PX / 2),
                            "color": ccol(self.theme.color(color_child(eff)))}
        font_color = None
        if style is not None and style.find(f"{A}fontRef") is not None:
            font_color = self.theme.color(color_child(style.find(f"{A}fontRef")))
        txbody = sp.find(f"{P}txBody")
        if txbody is not None:
            bps = [txbody.find(f"{A}bodyPr")] + [s.find(f"{P}txBody/{A}bodyPr") for s in ph_srcs]
            chain = [s.find(f"{P}txBody/{A}lstStyle") for s in ph_srcs]
            chain.append(self.master_text_style(info[0], ctx["master"]) if info else self.default_text)
            tx = self.text_of(txbody, chain, bps, font_color, ctx["no"])
            if tx:
                el["text"] = tx
        if "fill" not in el and "line" not in el and "text" not in el:
            return None
        return el

    @staticmethod
    def cust_path(geom, w, h):
        out = []
        for path in geom.findall(f"{A}pathLst/{A}path"):
            pw, ph_ = int(path.get("w", "0")) or None, int(path.get("h", "0")) or None
            sx = w / pw if pw else 1 / EMU_PX
            sy = h / ph_ if ph_ else 1 / EMU_PX
            cmds, cur = [], (0, 0)
            for cmd in path:
                t = lname(cmd)
                pts = [(r2(int(pt.get("x")) * sx), r2(int(pt.get("y")) * sy)) for pt in cmd.findall(f"{A}pt")]
                if t in ("moveTo", "lnTo"):
                    cmds.append(["M" if t == "moveTo" else "L", *pts[0]]); cur = pts[0]
                elif t == "cubicBezTo":
                    cmds.append(["C", *[v for p in pts for v in p]]); cur = pts[-1]
                elif t == "quadBezTo":
                    cmds.append(["Q", *[v for p in pts for v in p]]); cur = pts[-1]
                elif t == "close":
                    cmds.append(["Z"])
                elif t == "arcTo":
                    wr, hr = int(cmd.get("wR")) * sx, int(cmd.get("hR")) * sy
                    st, sw = int(cmd.get("stAng")) / 60000, int(cmd.get("swAng")) / 60000
                    cx = cur[0] - wr * math.cos(math.radians(st))
                    cy = cur[1] - hr * math.sin(math.radians(st))
                    for i in range(1, 17):
                        a = math.radians(st + sw * i / 16)
                        cur = (r2(cx + wr * math.cos(a)), r2(cy + hr * math.sin(a)))
                        cmds.append(["L", *cur])
            pj = {"cmds": cmds}
            if path.get("fill") == "none":
                pj["fill"] = False
            if path.get("stroke") == "0":
                pj["stroke"] = False
            out.append(pj)
        return out

    def picture(self, pic, ctx, xf):
        x = pic.find(f"{P}spPr/{A}xfrm")
        if x is None:
            return None
        X, Y, W, H = self.box_of(x, xf)
        el = {"x": X, "y": Y, "w": W, "h": H}
        rot = int(x.get("rot", "0")) / 60000
        if rot:
            el["rot"] = round(rot, 2)
        nm = pic.find(f".//{P}cNvPr")
        if nm is not None:
            if nm.get("name"):
                el["name"] = nm.get("name")
            if nm.get("descr"):
                el["alt"] = nm.get("descr")
        rels = ctx["rels"]
        svg_blip = pic.find(f".//{{{NS['asvg']}}}svgBlip")
        if svg_blip is not None:
            el["kind"] = "svg"
            el["svg"] = self.clean_svg(rels[svg_blip.get(f"{R}embed")].target_part.blob)
            return el
        blip = pic.find(f"{P}blipFill/{A}blip")
        rid = blip.get(f"{R}embed") if blip is not None else None
        if rid and W * H >= self.W * self.H * 0.8:
            grid = self.backdrop_grid(rels[rid].target_part.blob)
            if grid:
                el["kind"] = "backdrop"
                el["grid"] = grid
                return el
        el["kind"] = "placeholder"
        el["label"] = "画像（スクリーンショット）" + (f"：{el['alt']}" if el.get("alt") else "")
        return el

    @staticmethod
    def clean_svg(blob):
        root = _xml(blob)
        for e in list(root.iter()):
            if not isinstance(e.tag, str):
                continue
            if lname(e) in ("script", "foreignObject"):
                e.getparent().remove(e)
                continue
            for k in list(e.attrib):
                if k.lower().startswith("on"):
                    del e.attrib[k]
        for k in ("width", "height"):
            root.attrib.pop(k, None)
        s = etree.tostring(root, encoding="unicode")
        return re.sub(r">\s+<", "><", s)

    @staticmethod
    def backdrop_grid(blob, gx=24, gy=14):
        try:
            from PIL import Image
            im = Image.open(io.BytesIO(blob)).convert("RGB").resize((gx, gy), Image.BILINEAR)
        except Exception:  # noqa: BLE001
            return None
        return {"cols": gx, "rows": gy,
                "colors": ["#%02X%02X%02X" % im.getpixel((i, j)) for j in range(gy) for i in range(gx)]}

    def frame(self, fr, ctx, xf):
        x = fr.find(f"{P}xfrm")
        X, Y, W, H = self.box_of(x, xf)
        gd = fr.find(f"{A}graphic/{A}graphicData")
        uri = gd.get("uri", "")
        if uri.endswith("/table"):
            return self.table(gd.find(f"{A}tbl"), ctx, X, Y, W, H)
        if uri.endswith("/ole"):
            ole = gd.find(f".//{P}oleObj")
            if ole is not None and "Excel" in (ole.get("progId") or ""):
                try:
                    from pptx_import_excel import excel_table  # type: ignore
                    t = excel_table(ctx["rels"][ole.get(f"{R}id")].target_part.blob, X, Y, W, H)
                    if t:
                        return t
                except Exception as ex:  # noqa: BLE001
                    warn(f"スライド{ctx['no']}: 埋め込みExcelの表化に失敗: {ex!r}")
            return {"kind": "placeholder", "x": X, "y": Y, "w": W, "h": H, "label": "埋め込みオブジェクト"}
        return {"kind": "placeholder", "x": X, "y": Y, "w": W, "h": H, "label": "未対応オブジェクト"}

    def table(self, tbl, ctx, X, Y, W, H):
        tpr = tbl.find(f"{A}tblPr")
        sid = tpr.findtext(f"{A}tableStyleId") if tpr is not None else None
        style = self.tbl_styles.get(sid) if sid else None
        flags = {k: (tpr.get(k) in ("1", "true")) if tpr is not None else False
                 for k in ("firstRow", "lastRow", "firstCol", "lastCol", "bandRow", "bandCol")}
        cols = [px(g.get("w")) for g in tbl.findall(f"{A}tblGrid/{A}gridCol")]
        rows = tbl.findall(f"{A}tr")
        nrow, ncol = len(rows), len(cols)

        def parts_for(ri, ci):
            ps = ["wholeTbl"]
            if flags["bandRow"]:
                br = ri - (1 if flags["firstRow"] else 0)
                if br >= 0:
                    ps.append("band1H" if br % 2 == 0 else "band2H")
            if flags["bandCol"]:
                bc = ci - (1 if flags["firstCol"] else 0)
                if bc >= 0:
                    ps.append("band1V" if bc % 2 == 0 else "band2V")
            if flags["firstCol"] and ci == 0:
                ps.append("firstCol")
            if flags["lastCol"] and ci == ncol - 1:
                ps.append("lastCol")
            if flags["firstRow"] and ri == 0:
                ps.append("firstRow")
            if flags["lastRow"] and ri == nrow - 1:
                ps.append("lastRow")
            return ps

        def style_cell(ri, ci):
            fill, color, bold = None, None, False
            if style is None:
                return fill, color, bold
            for part in parts_for(ri, ci):
                el = style.find(A + part)
                if el is None:
                    continue
                f = el.find(f"{A}tcStyle/{A}fill/{A}solidFill")
                if f is not None:
                    fill = ccol(self.theme.color(color_child(f)))
                if el.find(f"{A}tcStyle/{A}fill/{A}noFill") is not None:
                    fill = None
                tx = el.find(f"{A}tcTxStyle")
                if tx is not None:
                    c = color_child(tx)
                    if c is not None:
                        color = self.theme.color(c)
                    if tx.get("b") == "on":
                        bold = True
            return fill, color, bold

        def style_border(ri, ci, side):
            if style is None:
                return None
            res = None
            outer = {"l": ci == 0, "r": ci == ncol - 1, "t": ri == 0, "b": ri == nrow - 1}[side]
            full = {"l": "left", "r": "right", "t": "top", "b": "bottom"}[side]
            for part in parts_for(ri, ci):
                el = style.find(f"{A}{part}/{A}tcStyle/{A}tcBdr")
                if el is None:
                    continue
                key = full if (outer or part != "wholeTbl") else ("insideV" if side in "lr" else "insideH")
                b = el.find(f"{A}{key}/{A}ln")
                if b is not None:
                    res = b
            return res

        def ln_json(ln):
            if ln is None:
                return None
            ls = Line()
            ls.width = 0.75
            ls.merge(ln, self.theme)
            return ls.to_json() or False

        out_rows = []
        for ri, tr in enumerate(rows):
            cells = []
            for ci, tc in enumerate(tr.findall(f"{A}tc")):
                if tc.get("hMerge") in ("1", "true") or tc.get("vMerge") in ("1", "true"):
                    cells.append({"merged": "h" if tc.get("hMerge") in ("1", "true") else "v"})
                    continue
                span, rspan = int(tc.get("gridSpan", "1")), int(tc.get("rowSpan", "1"))
                tcpr = tc.find(f"{A}tcPr")
                sfill, scolor, sbold = style_cell(ri, ci)
                fill = sfill
                if tcpr is not None:
                    for ch in tcpr:
                        t = lname(ch)
                        if t == "solidFill":
                            fill = ccol(self.theme.color(color_child(ch)))
                        elif t == "noFill":
                            fill = None
                        elif t == "gradFill":
                            g = self.grad(ch)["gradient"]["stops"]
                            fill = g[0][1] if g else None
                cell = {}
                if span > 1:
                    cell["span"] = span
                if rspan > 1:
                    cell["rowspan"] = rspan
                if fill is not None:
                    cell["fill"] = fill
                cell["pad"] = [px(tcpr.get(k, d)) if tcpr is not None else px(d)
                               for k, d in (("marL", 91440), ("marT", 45720), ("marR", 91440), ("marB", 45720))]
                anchor = tcpr.get("anchor", "t") if tcpr is not None else "t"
                if anchor != "t":
                    cell["anchor"] = anchor
                if tcpr is not None and tcpr.get("vert") in ("eaVert", "vert"):
                    cell["vert"] = True
                borders = {}
                for side, tag in (("l", "lnL"), ("r", "lnR"), ("t", "lnT"), ("b", "lnB")):
                    ln = tcpr.find(A + tag) if tcpr is not None else None
                    rr = ri + rspan - 1 if side == "b" else ri
                    cc = ci + span - 1 if side == "r" else ci
                    v = ln_json(ln) if ln is not None else ln_json(style_border(rr, cc, side))
                    if v is not None:
                        borders[side] = v or None
                if borders:
                    cell["borders"] = borders
                bp = etree.Element(A + "bodyPr", lIns="0", rIns="0", tIns="0", bIns="0")
                tx = self.text_of(tc.find(f"{A}txBody"), [self.default_text], [bp],
                                  scolor, ctx["no"], bold=sbold)
                cell["paragraphs"] = tx["paragraphs"] if tx else []
                cells.append(cell)
            out_rows.append({"h": px(tr.get("h")), "cells": cells})
        return {"kind": "table", "x": X, "y": Y, "w": W, "h": H, "cols": cols, "rows": out_rows}

    # ---- slide
    def background(self, els):
        for el in els:
            bg = el.find(f"{P}cSld/{P}bg")
            if bg is None:
                continue
            bgpr = bg.find(f"{P}bgPr")
            if bgpr is not None:
                f = self.fill_of(bgpr, None)
                if isinstance(f, dict) and "gradient" in f:
                    return f["gradient"]["stops"][0][1]
                return f if isinstance(f, str) else "#FFFFFF"
            ref = bg.find(f"{P}bgRef")
            if ref is not None:
                c = self.theme.color(color_child(ref))
                return ccol(c) if c else "#FFFFFF"
        return "#FFFFFF"

    def convert(self, only=None):
        slides = []
        total = len(self.prs.slides)
        for no, slide in enumerate(self.prs.slides, 1):
            if only and no not in only:
                continue
            layout = slide.slide_layout
            master = layout.slide_master
            cmap = dict(self.base_map)
            for el in (layout._element, slide._element):
                ov = el.find(f"{P}clrMapOvr/{A}overrideClrMapping")
                if ov is not None:
                    cmap = dict(ov.attrib)
            self.theme.clrmap = cmap
            ctx = {"no": no, "layout": layout._element, "master": master._element, "rels": slide.part.rels}
            elements = []
            show_master = layout._element.get("showMasterSp", "1") != "0" and slide._element.get("showMasterSp", "1") != "0"
            ident = (0, 0, 1, 1, 0, 0)
            if show_master:
                elements += self.walk(master._element.find(f"{P}cSld/{P}spTree"), dict(ctx, rels=master.part.rels, skip_ph=True), ident)
            if slide._element.get("showMasterSp", "1") != "0":
                elements += self.walk(layout._element.find(f"{P}cSld/{P}spTree"), dict(ctx, rels=layout.part.rels, skip_ph=True), ident)
            elements += self.walk(slide._element.find(f"{P}cSld/{P}spTree"), ctx, ident)
            title = ""
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = re.sub(r"\s+", " ", slide.shapes.title.text_frame.text.replace("\v", " ")).strip()
            s = {"type": "canvas", "title": title, "source_page": no,
                 "background": self.background([slide._element, layout._element, master._element]),
                 "elements": elements}
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                nt = slide.notes_slide.notes_text_frame.text.strip()
                if nt:
                    s["notes"] = nt
            slides.append(s)
            print(f"  {no}/{total}: {title[:40]}（要素 {len(elements)}）", file=sys.stderr)
        return slides


def main():
    for st in (sys.stdout, sys.stderr):
        try:
            st.reconfigure(encoding="utf-8", errors="replace")
        except Exception:  # noqa: BLE001
            pass
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    ap = argparse.ArgumentParser(description="PPTX → canvas 型 deck.json")
    ap.add_argument("src")
    ap.add_argument("--into", required=True, help="出力先のデッキフォルダ（deck.json を書き出す）")
    ap.add_argument("--id", help="meta.id（省略時はフォルダ名）")
    ap.add_argument("--title", help="meta.title（省略時は PPTX のファイル名）")
    ap.add_argument("--slides", nargs="*", type=int, help="取り込むスライド番号（省略=全部）")
    ap.add_argument("--confirm", action="store_true",
                    help="既存の deck.json を置き換えることを承認する（元は deck.json.bak に退避）")
    args = ap.parse_args()
    src = Path(args.src)
    deck_dir = Path(args.into)
    out = deck_dir / "deck.json"
    # 既存デッキは丸ごと置き換わるため、--confirm が無ければ書き込まずに終了する（archify_import と同じ承認ゲート）
    if out.exists() and not args.confirm:
        print(f"error: {out} が既にあります。pptx_import.py は deck.json を丸ごと置き換えます（追記ではありません）。\n"
              "  置き換えてよいか利用者に確認し、明示的な承認を得てから --confirm を付けて再実行してください"
              "（元のファイルは deck.json.bak に退避されます）。別のフォルダを --into に指定することもできます。",
              file=sys.stderr)
        sys.exit(1)
    if not src.is_file():
        print(f"error: {src} が見つかりません", file=sys.stderr)
        sys.exit(1)
    imp = Importer(src)
    slides = imp.convert(set(args.slides) if args.slides else None)
    deck = {"meta": {"id": args.id or deck_dir.name, "title": args.title or src.stem,
                     "source": src.name, "generator": "pptx_import.py"},
            "slides": slides}
    deck_dir.mkdir(parents=True, exist_ok=True)
    if out.exists():
        bak = deck_dir / "deck.json.bak"
        bak.write_bytes(out.read_bytes())
        print(f"既存の deck.json を {bak} に退避しました", file=sys.stderr)
    out.write_text(json.dumps(deck, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"deck.json: {out}（{len(slides)}枚、{out.stat().st_size // 1024} KB、警告 {len(WARNINGS)} 件）")


if __name__ == "__main__":
    main()
