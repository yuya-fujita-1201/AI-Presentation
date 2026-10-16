"""PPTX に埋め込まれた Excel（OLE）を canvas の table 要素にする。

表示範囲はブックの oleSize（PowerPoint 上に見えている範囲）を使い、列幅・行高を枠 (W, H) に合わせて縮尺する。
"""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

THEME_FALLBACK = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"]


def _color(c):
    if c is None:
        return None
    try:
        if c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8 and c.rgb != "00000000":
            return "#" + c.rgb[2:]
        if c.type == "theme" and c.theme is not None and c.theme < len(THEME_FALLBACK):
            base = THEME_FALLBACK[c.theme]
            tint = c.tint or 0
            r, g, b = (int(base[i:i + 2], 16) for i in (0, 2, 4))
            if tint > 0:
                r, g, b = (round(v + (255 - v) * tint) for v in (r, g, b))
            elif tint < 0:
                r, g, b = (round(v * (1 + tint)) for v in (r, g, b))
            return "#%02X%02X%02X" % (r, g, b)
        if c.type == "indexed" and c.indexed is not None:
            from openpyxl.styles.colors import COLOR_INDEX
            if c.indexed < len(COLOR_INDEX) and c.indexed not in (64, 65):
                return "#" + COLOR_INDEX[c.indexed][2:]
    except Exception:  # noqa: BLE001
        return None
    return None


def _fmt(v, nf):
    if v is None:
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        m = re.search(r"0\.(0+)", nf or "")
        if nf and "%" in nf:
            d = len(m.group(1)) if m else 0
            return f"{v * 100:.{d}f}%"
        if m:
            return f"{v:,.{len(m.group(1))}f}" if "," in nf else f"{v:.{len(m.group(1))}f}"
        if isinstance(v, float):
            return f"{v:g}"
        return str(v)
    return str(v)


def _border(side, scale):
    if side is None or side.style is None:
        return None
    w = {"thin": 0.75, "hair": 0.5, "medium": 1.5, "thick": 2.25, "dashed": 0.75, "dotted": 0.75}.get(side.style, 0.75)
    out = {"color": _color(side.color) or "#000000", "width": round(max(w * scale, 0.5), 2)}
    if side.style in ("dashed", "mediumDashed"):
        out["dash"] = "dash"
    elif side.style in ("dotted", "hair"):
        out["dash"] = "sysDot"
    return out


# 埋め込みブックは信用しない: 展開後サイズと表示範囲のセル数に上限を設ける（zip bomb・巨大範囲による DoS 対策）
MAX_UNZIPPED = 50 * 1024 * 1024
MAX_ROWS, MAX_COLS, MAX_CELLS = 300, 60, 6000


def excel_table(blob, X, Y, W, H):
    import zipfile
    zf = zipfile.ZipFile(io.BytesIO(blob))
    if sum(i.file_size for i in zf.infolist()) > MAX_UNZIPPED:
        raise ValueError("埋め込みブックの展開後サイズが上限（50MB）を超えています")
    wbxml = zf.read("xl/workbook.xml").decode("utf-8", "replace")
    wb = load_workbook(io.BytesIO(blob), data_only=True)
    m = re.search(r'<oleSize ref="([A-Z]+\d+(?::[A-Z]+\d+)?)"', wbxml)
    ws = wb.worksheets[wb.index(wb.active)] if wb.active is not None else wb.worksheets[0]
    ref = m.group(1) if m else ws.dimensions
    if ":" not in ref:
        ref = f"{ref}:{ref}"
    c1, r1, c2, r2 = range_boundaries(ref)
    if r2 - r1 + 1 > MAX_ROWS or c2 - c1 + 1 > MAX_COLS or (r2 - r1 + 1) * (c2 - c1 + 1) > MAX_CELLS:
        raise ValueError(f"埋め込みブックの表示範囲 {ref} が大きすぎます（上限 {MAX_ROWS}行 × {MAX_COLS}列・{MAX_CELLS}セル）")
    widths = []
    for c in range(c1, c2 + 1):
        dim = ws.column_dimensions.get(get_column_letter(c))
        if dim is not None and dim.hidden:
            widths.append(0)
            continue
        wch = dim.width if dim is not None and dim.width else (ws.sheet_format.defaultColWidth or 8.43)
        widths.append(wch * 7 + 5)
    heights = []
    for r in range(r1, r2 + 1):
        dim = ws.row_dimensions.get(r)
        if dim is not None and dim.hidden:
            heights.append(0)
            continue
        hpt = dim.height if dim is not None and dim.height else (ws.sheet_format.defaultRowHeight or 15)
        heights.append(hpt * 96 / 72)
    sx = W / max(sum(widths), 1)
    sy = H / max(sum(heights), 1)
    merged = {}
    covered = {}
    for mr in ws.merged_cells.ranges:
        if mr.min_row > r2 or mr.max_row < r1 or mr.min_col > c2 or mr.max_col < c1:
            continue
        merged[(mr.min_row, mr.min_col)] = (min(mr.max_row, r2) - mr.min_row + 1, min(mr.max_col, c2) - mr.min_col + 1)
        for rr in range(mr.min_row, min(mr.max_row, r2) + 1):
            for cc in range(mr.min_col, min(mr.max_col, c2) + 1):
                if (rr, cc) != (mr.min_row, mr.min_col):
                    covered[(rr, cc)] = "h" if rr == mr.min_row else "v"
    rows = []
    for ri, r in enumerate(range(r1, r2 + 1)):
        cells = []
        for c in range(c1, c2 + 1):
            if (r, c) in covered:
                cells.append({"merged": covered[(r, c)]})
                continue
            cell = ws.cell(row=r, column=c)
            fs = (cell.font.sz or 11) * 96 / 72 * sy if cell.font else 11 * 96 / 72 * sy
            run = {"text": _fmt(cell.value, cell.number_format), "size": round(fs, 2),
                   "color": _color(cell.font.color) if cell.font else None, "font": [cell.font.name or "Meiryo UI", "Meiryo UI"]}
            if run["color"] is None:
                run["color"] = "#000000"
            if cell.font and cell.font.b:
                run["bold"] = True
            al = cell.alignment.horizontal if cell.alignment else None
            if al is None and isinstance(cell.value, (int, float)):
                al = "right"
            para = {"runs": [run] if run["text"] else [], "lineHeight": 1.0}
            if not run["text"]:
                para["size"] = run["size"]
            if al in ("center", "centerContinuous"):
                para["align"] = "ctr"
            elif al == "right":
                para["align"] = "r"
            jc = {"top": "t", "center": "ctr", "bottom": "b"}.get(cell.alignment.vertical if cell.alignment else None, "b")
            # Excel は折り返し指定のないセルを折り返さない（はみ出し分は隣のセルに隠れる）
            out = {"pad": [2, 0, 2, 0], "anchor": jc, "paragraphs": [para]}
            if not (cell.alignment and cell.alignment.wrap_text):
                out["nowrap"] = True
                nxt = ws.cell(row=r, column=c + 1).value if c < c2 else 0
                # 数値は隣へはみ出さない。文字列も右隣が空でなければ切れる
                if isinstance(cell.value, (int, float)) or (nxt not in (None, "")):
                    out["clip"] = True
            fill = cell.fill
            if fill is not None and fill.fill_type == "solid":
                fc = _color(fill.fgColor)
                if fc:
                    out["fill"] = fc
            b = {}
            for k, side in (("l", cell.border.left), ("r", cell.border.right), ("t", cell.border.top), ("b", cell.border.bottom)):
                v = _border(side, min(sx, sy) * 1.3)
                if v:
                    b[k] = v
            if b:
                out["borders"] = b
            if (r, c) in merged:
                rs, cs = merged[(r, c)]
                if cs > 1:
                    out["span"] = cs
                if rs > 1:
                    out["rowspan"] = rs
            cells.append(out)
        rows.append({"h": round(heights[ri] * sy, 2), "cells": cells})
    return {"kind": "table", "x": X, "y": Y, "w": W, "h": H, "cols": [round(w * sx, 2) for w in widths],
            "rows": rows, "name": "埋め込みExcel"}
