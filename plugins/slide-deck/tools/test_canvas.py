#!/usr/bin/env python3
"""canvas 型（canvas_engine.py）と PPTX 取り込み（pptx_import.py）の回帰テスト（unittest）。

    python -m unittest discover -s tools -p "test_*.py"
"""
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))
import canvas_engine as ce  # noqa: E402

PY = sys.executable


def _deck(slides):
    return {"meta": {"id": "t", "title": "t"}, "slides": slides}


SHAPES = [
    {"x": 10, "y": 20, "w": 200, "h": 40, "geom": "homePlate", "adj": {"adj": 13400}, "fill": "#F0E7FF",
     "line": {"color": "#7F7F7F", "width": 0.75},
     "text": {"anchor": "ctr", "paragraphs": [{"align": "ctr", "runs": [{"text": "要件定義", "size": 12}]}]}},
    {"x": 10, "y": 80, "w": 100, "h": 0, "geom": "straightConnector1",
     "line": {"color": "#000000", "width": 0.75, "dash": "dash", "tail": {"type": "triangle"}}},
    {"x": 300, "y": 20, "w": 100, "h": 40, "fill": {"pattern": "ltDnDiag", "fg": "#A100FF", "bg": "#FFFFFF"}},
    {"kind": "table", "x": 10, "y": 200, "w": 200, "h": 40, "cols": [100, 100],
     "rows": [{"h": 20, "cells": [{"fill": "#A100FF", "paragraphs": [{"runs": [{"text": "見出し", "size": 12, "bold": True, "color": "#FFFFFF"}]}]},
                                  {"paragraphs": [{"runs": [{"text": "値", "size": 12}]}], "borders": {"b": {"color": "#000000", "width": 4}}}]},
              {"h": 20, "cells": [{"span": 2, "paragraphs": [{"runs": [{"text": "結合", "size": 12}]}]}, {"merged": "h"}]}]},
    {"kind": "svg", "x": 500, "y": 20, "w": 40, "h": 40, "svg": '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><defs><linearGradient id="a"/></defs><rect width="10" height="10" fill="url(#a)" onclick="x()"/></svg>'},
    {"kind": "placeholder", "x": 600, "y": 20, "w": 100, "h": 60, "label": "画像（スクリーンショット）"},
]


class CanvasHtmlTest(unittest.TestCase):
    def test_shapes_render(self):
        html = ce.canvas_html({"elements": SHAPES})
        self.assertIn("要件定義", html)
        self.assertIn("<pattern", html)
        self.assertIn("stroke-dasharray", html)
        self.assertIn('colspan="2"', html)
        self.assertNotIn("onclick", html)  # SVG のイベント属性は除去する
        self.assertNotIn('id="a"', html)   # SVG の id は接頭辞付きに置き換える
        self.assertIn("画像（スクリーンショット）", html)

    def test_homeplate_text_rect_is_narrower(self):
        l, t, r, b = ce.text_rect({"geom": "homePlate", "w": 200, "h": 40, "adj": {"adj": 50000}})
        self.assertLess(r, 200)
        l, t, r, b = ce.text_rect({"geom": "flowChartDecision", "w": 200, "h": 40})
        self.assertEqual((l, t, r, b), (50, 10, 150, 30))

    def test_mixed_size_lines_are_split(self):
        p = {"runs": [{"text": "75.1", "size": 24}, {"br": True}, {"text": "(▲6.6)", "size": 12}]}
        parts = ce._split_lines(p)
        self.assertEqual(len(parts), 2)
        self.assertEqual(parts[1]["runs"][0]["text"], "(▲6.6)")

    def test_table_border_compensated(self):
        html = ce.table_html(SHAPES[3])
        # 4px の下罫線の半分（2px）だけ下余白が減る（既定 4.8px → 2.8px）
        self.assertIn("padding:4.8px 9.6px 2.8px 9.6px", html)

    def test_validate_rejects_unknown_kind_and_script(self):
        errors = []

        def fail(msg):
            errors.append(msg)
            raise SystemExit(1)
        with self.assertRaises(SystemExit):
            ce.validate_canvas({"elements": [{"kind": "video", "x": 0, "y": 0}]}, 0, fail, lambda m: None)
        with self.assertRaises(SystemExit):
            ce.validate_canvas({"elements": [{"kind": "svg", "x": 0, "y": 0, "svg": "<svg><script>1</script></svg>"}]},
                               0, fail, lambda m: None)
        self.assertEqual(len(errors), 2)


class CanvasSanitizeTest(unittest.TestCase):
    """deck.json の値が HTML 属性・CSS・PPTX の XML へ素通しされないこと（取り込み元は信用しない）"""

    def _issues(self, els):
        issues = []
        ce.sanitize_elements(els, "elements", issues)
        return issues

    def test_svg_allowlist(self):
        s = ce.sanitize_svg(
            '<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" onload="x()">'
            '<script>alert(1)</script><foreignObject><div/></foreignObject>'
            '<a href="javascript:alert(1)"><rect/></a><image href="https://evil.example/x.png"/>'
            '<use xlink:href="#p"/><path id="p" d="M0 0" fill="url(#g)" style="fill:url(https://evil.example)"/></svg>')
        for bad in ("onload", "script", "foreignObject", "javascript", "evil", "<a ", "<image"):
            self.assertNotIn(bad, s)
        self.assertIn('<use href="#p"', s)
        self.assertIn('fill="url(#g)"', s)
        self.assertEqual(ce.sanitize_svg('<!DOCTYPE svg [<!ENTITY x "y">]><svg>&x;</svg>'), "")
        self.assertEqual(ce.sanitize_svg("<svg><rect></svg>"), "")
        self.assertEqual(ce.sanitize_svg("<div/>"), "")

    def test_bad_values_are_reported_and_neutralized(self):
        els = [{"x": 0, "y": 0, "w": 10, "h": 10, "fill": 'red"/><x onload="1',
                "line": {"color": "#000000", "dash": 'x"', "tail": {"type": 'x"', "w": "huge"}},
                "text": {"anchor": 'ctr" foo="', "paragraphs": [{"align": "zz", "runs": [
                    {"text": "a", "color": "#12345", "font": ["Meiryo UI'\"><b>"]}]}]}},
               {"geom": "custom", "x": 0, "y": 0, "w": 1, "h": 1, "path": [{"cmds": [['M"/>', 0, 0], ["L", 1, 1]]}]},
               {"kind": "table", "x": 0, "y": 0, "cols": [10], "rows": [{"h": 10, "cells": [{"span": '2" evilspan="'}]}]}]
        issues = self._issues(els)
        self.assertGreaterEqual(len(issues), 8, issues)
        html = ce.canvas_html({"elements": els})
        for bad in ("onload", 'foo="', "<b>", "evilspan", 'M"'):
            self.assertNotIn(bad, html)
        from pptx import Presentation
        prs = Presentation()
        ce.add_canvas_to_pptx(prs.slides.add_slide(prs.slide_layouts[6]), els)  # 不正な XML なら parse_xml で例外

    def test_imported_values_pass(self):
        self.assertEqual(self._issues([e for e in SHAPES if e.get("kind") != "svg"]), [])

    def test_validate_fails_on_bad_color(self):
        with self.assertRaises(SystemExit):
            ce.validate_canvas({"elements": [{"x": 0, "y": 0, "w": 1, "h": 1, "fill": "purple"}]}, 0,
                               lambda m: (_ for _ in ()).throw(SystemExit(1)), lambda m: None)


class CanvasBuildTest(unittest.TestCase):
    def test_build_html_and_pptx(self):
        with tempfile.TemporaryDirectory() as d:
            deck_dir = Path(d) / "t"
            deck_dir.mkdir()
            els = [e for e in SHAPES if e.get("kind") != "svg"]  # PPTX の SVG 画像化は Playwright 依存のため除外
            (deck_dir / "deck.json").write_text(json.dumps(_deck([
                {"type": "canvas", "title": "T", "background": "#FFFFFF", "elements": els},
                {"type": "bullets", "title": "通常スライド", "bullets": ["a"]}]), ensure_ascii=False), encoding="utf-8")
            r = subprocess.run([PY, str(TOOLS / "build_deck.py"), str(deck_dir)], capture_output=True, text=True,
                               encoding="utf-8")
            self.assertEqual(r.returncode, 0, r.stderr)
            html = (deck_dir / "build" / "t.html").read_text(encoding="utf-8")
            self.assertIn('data-type="canvas"', html)
            self.assertIn(".cv-s{", html)
            from pptx import Presentation
            prs = Presentation(str(deck_dir / "build" / "t.pptx"))
            texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
            self.assertIn("要件定義", texts)

    def test_bad_background_fails(self):
        with tempfile.TemporaryDirectory() as d:
            deck_dir = Path(d) / "t"
            deck_dir.mkdir()
            (deck_dir / "deck.json").write_text(json.dumps(_deck([
                {"type": "canvas", "background": "purple", "elements": []}])), encoding="utf-8")
            r = subprocess.run([PY, str(TOOLS / "build_deck.py"), str(deck_dir), "--html"], capture_output=True,
                               text=True, encoding="utf-8")
            self.assertNotEqual(r.returncode, 0)


class PptxImportTest(unittest.TestCase):
    def test_roundtrip(self):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Emu
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "src.pptx"
            prs = Presentation()
            prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
            s = prs.slides.add_slide(prs.slide_layouts[6])
            sh = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Emu(952500), Emu(952500), Emu(1905000), Emu(476250))
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor(0xA1, 0x00, 0xFF)
            sh.text_frame.text = "Plan"
            tb = s.shapes.add_table(2, 2, Emu(952500), Emu(2857500), Emu(3810000), Emu(762000)).table
            tb.cell(0, 0).text = "見出し"
            tb.cell(1, 1).text = "値"
            prs.save(str(src))
            deck_dir = Path(d) / "deck"
            r = subprocess.run([PY, str(TOOLS / "pptx_import.py"), str(src), "--into", str(deck_dir)],
                               capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, r.stderr)
            deck = json.loads((deck_dir / "deck.json").read_text(encoding="utf-8"))
            els = deck["slides"][0]["elements"]
            chev = next(e for e in els if e.get("geom") == "chevron")
            self.assertEqual((chev["x"], chev["y"], chev["w"], chev["h"]), (100, 100, 200, 50))
            self.assertEqual(chev["fill"], "#A100FF")
            self.assertEqual(chev["text"]["paragraphs"][0]["runs"][0]["text"], "Plan")
            table = next(e for e in els if e.get("kind") == "table")
            self.assertEqual(table["rows"][0]["cells"][0]["paragraphs"][0]["runs"][0]["text"], "見出し")
            # 既定の表スタイル（中間スタイル2-アクセント1）の見出し行は太字・白文字になる
            run = table["rows"][0]["cells"][0]["paragraphs"][0]["runs"][0]
            self.assertTrue(run.get("bold"))

            # 既存 deck.json があると --confirm 無しでは書き込まない（丸ごと置き換えになるため）
            before = (deck_dir / "deck.json").read_bytes()
            r = subprocess.run([PY, str(TOOLS / "pptx_import.py"), str(src), "--into", str(deck_dir)],
                               capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 1)
            self.assertIn("--confirm", r.stderr)
            self.assertEqual((deck_dir / "deck.json").read_bytes(), before)
            self.assertFalse((deck_dir / "deck.json.bak").exists())
            r = subprocess.run([PY, str(TOOLS / "pptx_import.py"), str(src), "--into", str(deck_dir), "--confirm"],
                               capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertEqual((deck_dir / "deck.json.bak").read_bytes(), before)


class PptxImportExcelTest(unittest.TestCase):
    def test_huge_range_is_rejected(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl 未導入")
        import io
        from pptx_import_excel import excel_table
        wb = openpyxl.Workbook()
        wb.active["A1"] = "x"
        wb.active["ZZ5000"] = "y"  # oleSize が無いので ws.dimensions（A1:ZZ5000）が表示範囲になる
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaises(ValueError):
            excel_table(buf.getvalue(), 0, 0, 100, 100)


if __name__ == "__main__":
    unittest.main()
