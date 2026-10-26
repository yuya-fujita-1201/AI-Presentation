"""proposal_engine（summary / reasons / statement / row_matrix / org と共通パーツ）のテスト。"""
import json
import sys
import tempfile
import unittest
from pathlib import Path

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))
import build_deck as b  # noqa: E402
import canvas_engine as ce  # noqa: E402
import proposal_engine as pe  # noqa: E402

LAYOUT = b.load_json(b.ROOT / "templates" / "layouts" / "default.json")
THEME = b.load_theme("default", b.theme_dirs(None))

SLIDES = {
    "summary": {"type": "summary", "title": "要旨", "lead": "リード",
                "items": [{"lead": "背景｜A", "text": "説明"}, {"lead": "課題｜B", "text": "説明"}, "要点だけ"],
                "source": "出典: X", "conclusion": "結論"},
    "reasons": {"type": "reasons", "title": "理由",
                "issues": [{"title": "課題1", "text": "t"}, "課題2"],
                "strengths": [{"title": "強み1", "items": ["根拠"]}, {"title": "強み2", "label": "B"}]},
    "statement": {"type": "statement", "kicker": "結論", "message": "大きな一文", "support": "補足"},
    "row_matrix": {"type": "row_matrix", "title": "整理", "header_style": "arrow", "emph_col": 1,
                   "columns": ["A", "B"], "rows": [{"label": "行1", "cells": [["x", "y"], {"text": "t", "items": ["i"]}]},
                                                   {"label": "行2", "cells": ["s", "u"]}]},
    "org": {"type": "org", "title": "体制",
            "left": {"title": "貴社", "nodes": [{"id": "a", "title": "A"}, {"id": "b", "parent": "a", "title": "B"}]},
            "right": {"title": "自社", "scope": True, "nodes": [{"id": "c", "title": "C", "name": "1名"}]},
            "links": [{"from": "a", "to": "c", "label": "連携"}]},
}


def elements(slide):
    st = b.resolve_style(LAYOUT, {"meta": {}}, slide)
    return pe.layout(slide, st, THEME)


class LayoutTests(unittest.TestCase):
    def test_all_types_produce_safe_elements_inside_canvas(self):
        for t, slide in SLIDES.items():
            with self.subTest(t=t):
                els = elements(slide)
                self.assertTrue(els)
                issues = []
                ce.sanitize_elements(els, issues=issues)
                self.assertEqual(issues, [], issues)   # 色は hex に解決され、canvas の検査をそのまま通る
                for e in els:
                    self.assertGreaterEqual(e["x"], -1)
                    self.assertGreaterEqual(e["y"], -1)
                    self.assertLessEqual(e["x"] + e["w"], 1281)
                    self.assertLessEqual(e["y"] + e["h"], 721)

    def test_common_parts_reserve_space(self):
        els = elements(SLIDES["summary"])
        names = {e.get("name"): e for e in els}
        self.assertIn("結論", names)
        self.assertIn("出典", names)
        last_row = max(e["y"] + e["h"] for e in els if str(e.get("name", "")).startswith("要点"))
        self.assertLessEqual(last_row, names["結論"]["y"])
        self.assertLessEqual(names["結論"]["y"] + names["結論"]["h"], names["出典"]["y"])

    def test_summary_rows_equal_height_and_numbered(self):
        els = elements(SLIDES["summary"])
        rows = [e for e in els if str(e.get("name", "")).startswith("要点")]
        self.assertEqual(len(rows), 3)
        self.assertEqual(len({e["h"] for e in rows}), 1)
        nums = [e["text"]["paragraphs"][0]["runs"][0]["text"] for e in els if str(e.get("name", "")).startswith("番号")]
        self.assertEqual(nums, ["1", "2", "3"])

    def test_long_text_shrinks_font(self):
        short = elements(SLIDES["summary"])
        long_slide = dict(SLIDES["summary"], items=[{"lead": "長い要点" * 6, "text": "長い説明の文章です。" * 12}] * 5)
        long_ = elements(long_slide)

        def lead_size(els):
            return next(e for e in els if e.get("name") == "要点1")["text"]["paragraphs"][0]["runs"][0]["size"]
        self.assertLess(lead_size(long_), lead_size(short))

    def test_reasons_labels_and_triangles(self):
        els = elements(SLIDES["reasons"])
        labels = [e["text"]["paragraphs"][0]["runs"][0]["text"] for e in els if str(e.get("name", "")).startswith("強み記号")]
        self.assertEqual(labels, ["a", "B"])
        self.assertEqual(sum(1 for e in els if str(e.get("name", "")).startswith("対応")), 2)

    def test_statement_has_no_chrome_and_full_background(self):
        els = elements(SLIDES["statement"])
        self.assertEqual((els[0]["w"], els[0]["h"]), (1280, 720))
        self.assertEqual(els[0]["fill"], THEME["colors"]["primary"].upper())

    def test_row_matrix_arrow_headers_and_emphasis(self):
        els = elements(SLIDES["row_matrix"])
        heads = [e for e in els if str(e.get("name", "")).startswith("列見出し")]
        self.assertEqual([e["geom"] for e in heads], ["homePlate", "chevron"])
        cells = {e["name"]: e for e in els if str(e.get("name", "")).startswith("セル")}
        self.assertNotEqual(cells["セル1-1"]["fill"], cells["セル1-2"]["fill"])

    def test_org_tree_and_scope(self):
        els = elements(SLIDES["org"])
        names = {e.get("name") for e in els}
        self.assertTrue({"a", "b", "c", "提案範囲", "提案範囲の札", "連携の名前"} <= names)
        a = next(e for e in els if e.get("name") == "a")
        bb = next(e for e in els if e.get("name") == "b")
        self.assertGreater(bb["y"], a["y"])   # 子は親より下の段


class ValidateTests(unittest.TestCase):
    def _errors(self, slide):
        errs, warns = [], []
        pe.validate(slide, 0, lambda m: errs.append(m) or (_ for _ in ()).throw(SystemExit(1)), warns.append)
        return errs, warns

    def test_schema_errors(self):
        cases = [{"type": "summary"}, {"type": "reasons", "issues": []}, {"type": "statement"},
                 {"type": "row_matrix", "rows": []},
                 {"type": "row_matrix", "columns": ["a", "b"], "col_ratios": [1], "rows": [{"label": "x"}]},
                 {"type": "org"},
                 {"type": "org", "left": {"nodes": [{"id": "a"}, {"id": "a"}]}},
                 {"type": "org", "left": {"nodes": [{"id": "a", "parent": "zz"}]}},
                 {"type": "org", "left": {"nodes": [{"id": "a"}]}, "links": [{"from": "a", "to": "q"}]}]
        for c in cases:
            with self.subTest(c=c):
                with self.assertRaises(SystemExit):
                    self._errors(c)

    def test_warnings_for_crowded_slides(self):
        _, w = self._errors({"type": "summary", "items": ["x"] * 8})
        self.assertTrue(w)
        _, w = self._errors(SLIDES["org"])
        self.assertEqual(w, [])


class BuildTests(unittest.TestCase):
    def test_build_html_and_pptx(self):
        from pptx import Presentation
        deck = {"meta": {"title": "t"}, "slides": list(SLIDES.values())}
        with tempfile.TemporaryDirectory() as d:
            dd = Path(d)
            (dd / "deck.json").write_text(json.dumps(deck, ensure_ascii=False), encoding="utf-8")
            deck_, theme, layout = b.load_deck(dd)
            b.validate_deck(deck_, layout, theme, dd)
            ex = b.expand_slides(deck_, layout)
            b.build_html(ex, theme, layout, dd, dd / "o.html")
            b.build_pptx(ex, theme, layout, dd, dd / "o.pptx")
            html = (dd / "o.html").read_text(encoding="utf-8")
            self.assertIn("cv-s", html)
            self.assertIn("大きな一文", html)
            prs = Presentation(str(dd / "o.pptx"))
            self.assertEqual(len(prs.slides), 5)
            for s in prs.slides:
                self.assertFalse(s.has_notes_slide)
                ids = [sh.shape_id for sh in s.shapes]
                self.assertEqual(len(ids), len(set(ids)))
            texts = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
            self.assertIn("背景｜A", texts)
            self.assertIn("結論", texts)


if __name__ == "__main__":
    unittest.main()
