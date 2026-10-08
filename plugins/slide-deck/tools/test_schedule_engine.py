#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""schedule_engine.py（矢羽根スケジュール）と schedule タイプのビルドのテスト。"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

TOOLS_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS_DIR))
import build_deck  # noqa: E402
import schedule_engine as se  # noqa: E402

ROOT = TOOLS_DIR.parent
LAYOUT = build_deck.load_json(ROOT / "templates" / "layouts" / "default.json")
CHART = LAYOUT["types"]["schedule"]["chart"]


def _slide(**kw):
    s = {"type": "schedule", "title": "t",
         "timeline": {"start": "2026-09", "end": "2027-02", "unit": "month"},
         "rows": [{"label": "A", "bars": [{"start": "2026-09", "end": "2026-10", "label": "x"}]}]}
    s.update(kw)
    return s


def _codes(diags, level=None):
    return [d["code"] for d in diags if level is None or d["level"] == level]


class AxisTest(unittest.TestCase):
    def test_month_timeline_labels_and_positions(self):
        ax = se.Axis(_slide(), [])
        self.assertEqual(["9月", "10月", "11月", "12月", "1月", "2月"], ax.labels)
        self.assertEqual(0.0, ax.pos("2026-09", "start"))
        self.assertEqual(2.0, ax.pos("2026-10", "end"))
        self.assertAlmostEqual(1 + 30 / 31, ax.pos("2026-10-30", "end"))
        self.assertAlmostEqual(1 + 29 / 31, ax.pos("2026-10-30", "start"))
        self.assertEqual(4.0, ax.pos("2027-01", "start"))  # 年またぎ

    def test_column_label_and_number_positions(self):
        ax = se.Axis(_slide(), [])
        self.assertEqual(1.0, ax.pos("10月", "start"))
        self.assertEqual(2.0, ax.pos("10月", "end"))
        self.assertEqual(1.5, ax.pos("10月", "at"))
        self.assertEqual(2.25, ax.pos(2.25, "start"))
        self.assertIsNone(ax.pos("来月", "start"))

    def test_week_timeline(self):
        ax = se.Axis(_slide(timeline={"start": "2026-10-05", "end": "2026-11-02", "unit": "week"}), [])
        self.assertEqual(["10/5", "10/12", "10/19", "10/26", "11/2"], ax.labels)
        self.assertEqual(1.0, ax.pos("2026-10-12", "start"))
        self.assertAlmostEqual(3.0, ax.pos("2026-10-25", "end"))  # 10/19〜10/25 の週の右端

    def test_explicit_columns_without_timeline(self):
        ax = se.Axis({"columns": ["Q1", "Q2", "Q3"]}, [])
        self.assertEqual(3, ax.n)
        self.assertEqual(3.0, ax.pos("Q3", "end"))


class ValidateTest(unittest.TestCase):
    def test_clean_slide_has_no_errors(self):
        self.assertEqual([], _codes(se.validate_schedule(_slide()), "error"))

    def test_missing_axis_is_error(self):
        s = _slide(); del s["timeline"]
        self.assertIn("missing-axis", _codes(se.validate_schedule(s), "error"))

    def test_out_of_range_and_reversed_bars_are_errors(self):
        s = _slide(rows=[{"label": "A", "bars": [
            {"start": "2026-08", "end": "2026-10"},
            {"start": "2026-12", "end": "2026-11"}]}])
        codes = _codes(se.validate_schedule(s), "error")
        self.assertIn("out-of-range", codes)
        self.assertIn("empty-bar", codes)

    def test_unparseable_position_is_error(self):
        s = _slide(rows=[{"label": "A", "bars": [{"start": "そのうち", "end": "2026-10"}]}])
        self.assertIn("invalid-position", _codes(se.validate_schedule(s), "error"))

    def test_overlapping_bars_and_unknown_variant_warn(self):
        s = _slide(rows=[{"label": "A", "bars": [
            {"start": "2026-09", "end": "2026-11"},
            {"start": "2026-10", "end": "2026-12", "variant": "purple"}]}])
        codes = _codes(se.validate_schedule(s), "warning")
        self.assertIn("bar-overlap", codes)
        self.assertIn("unknown-variant", codes)

    def test_periods_span_over_columns_is_error(self):
        s = _slide(periods=[{"label": "H1", "span": 4}, {"label": "H2", "span": 4}])
        self.assertIn("invalid-periods", _codes(se.validate_schedule(s), "error"))


class LayoutTest(unittest.TestCase):
    def test_bars_align_to_time_axis(self):
        geo = se.layout_schedule(_slide(), CHART)
        b = geo["bars"][0]
        gap = CHART["bar_gap"] / 2
        self.assertAlmostEqual(geo["tx0"] + gap, b["x"])
        self.assertAlmostEqual(geo["tx0"] + 2 * geo["col_w"] - gap, b["x"] + b["w"])
        self.assertEqual("arrow", b["shape"])

    def test_everything_stays_inside_chart_region(self):
        s = _slide(today="2026-10-05", note="※ 注記",
                   rows=[{"group": "G", "label": "節目", "milestones": [
                       {"at": "2026-10-31", "label": "10月末 レビュー完了"},
                       {"at": "2026-11-05", "label": "11/5 承認"}]}]
                   + [{"group": "G", "label": f"行{i}", "bars": [{"start": "2026-09", "end": "2027-02", "label": "全期間"}]}
                      for i in range(10)])
        geo = se.layout_schedule(s, CHART)
        rg = geo["region"]
        bottom = rg["y"] + rg["h"]
        for b in geo["bars"]:
            self.assertGreaterEqual(b["x"], rg["x"])
            self.assertLessEqual(b["x"] + b["w"], rg["x"] + rg["w"] + 0.01)
        self.assertLessEqual(geo["rows"][-1]["y"] + geo["rows"][-1]["h"], bottom)
        for it in geo["legend"]["items"]:
            self.assertLessEqual(it["y"] + it["sym_h"], bottom + 2)
        # 近接した節目ラベルは別の段に積まれる
        ms = geo["milestones"]
        self.assertEqual(2, len(ms))
        self.assertNotEqual(ms[0]["ly"], ms[1]["ly"])
        self.assertIsNotNone(geo["today"])
        self.assertEqual(["main", "milestone", "today"],
                         [it.get("variant", it["kind"]) for it in geo["legend"]["items"]])

    def test_long_label_in_short_bar_warns(self):
        s = _slide(rows=[{"label": "A", "bars": [
            {"start": "2026-09-01", "end": "2026-09-05", "label": "とても長いラベルが入っています"}]}])
        self.assertIn("bar-label-overflow", _codes(se.layout_schedule(s, CHART)["diagnostics"], "warning"))

    def test_label_pos_after_skips_overflow_warning(self):
        s = _slide(rows=[{"label": "A", "bars": [
            {"start": "2026-09-01", "end": "2026-09-05", "label": "とても長いラベル", "label_pos": "after"}]}])
        self.assertNotIn("bar-label-overflow", _codes(se.layout_schedule(s, CHART)["diagnostics"]))

    def test_legend_can_be_disabled(self):
        self.assertIsNone(se.layout_schedule(_slide(legend=False), CHART)["legend"])


class BuildTest(unittest.TestCase):
    def _build(self, slide):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        deck_dir = Path(tmp.name) / "sched"
        deck_dir.mkdir()
        (deck_dir / "deck.json").write_text(json.dumps(
            {"meta": {"title": "s"}, "slides": [slide]}, ensure_ascii=False), encoding="utf-8")
        proc = subprocess.run([sys.executable, str(TOOLS_DIR / "build_deck.py"), str(deck_dir)],
                              capture_output=True, text=True, encoding="utf-8", errors="replace")
        return deck_dir, proc

    def test_pptx_uses_editable_native_shapes(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
        slide = _slide(today="2026-10-05", rows=[
            {"label": "節目", "milestones": [{"at": "2026-10-31", "label": "RV完了"}]},
            {"label": "A", "bars": [{"start": "2026-09", "end": "2026-10", "label": "要件定義"},
                                    {"start": "2026-11", "end": "2026-12", "label": "設計", "shape": "chevron",
                                     "variant": "tbd"}]}])
        deck_dir, proc = self._build(slide)
        self.assertEqual(0, proc.returncode, proc.stderr)
        prs = Presentation(deck_dir / "build" / "sched.pptx")
        self.assertFalse(prs.slides[0].has_notes_slide)
        found = {}

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(sh.shapes)
                elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.has_text_frame:
                    found.setdefault(sh.auto_shape_type, []).append(sh.text_frame.text)

        walk(prs.slides[0].shapes)
        self.assertIn("要件定義", found.get(MSO_SHAPE.PENTAGON, []))
        self.assertIn("設計", found.get(MSO_SHAPE.CHEVRON, []))
        self.assertIn(MSO_SHAPE.DIAMOND, found)
        html = (deck_dir / "build" / "sched.html").read_text(encoding="utf-8")
        self.assertIn("要件定義", html)
        self.assertIn("<polygon", html)

    def test_schema_error_stops_build(self):
        _deck_dir, proc = self._build(_slide(rows=[{"label": "A", "bars": [{"start": "2025-01", "end": "2026-10"}]}]))
        self.assertNotEqual(0, proc.returncode)
        self.assertIn("out-of-range", proc.stderr)


if __name__ == "__main__":
    unittest.main()
