"""ja_text（字幅の実測・折り返しの模擬・泣き別れ対策）と build_deck への組み込みのテスト。"""
import sys
import unittest
from pathlib import Path
from unittest import mock

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))
import ja_text  # noqa: E402
import build_deck as b  # noqa: E402


def fallback_only():
    """フォントを実測しない（近似だけで計算する）状態にする。"""
    return mock.patch.object(ja_text, "_pick_font", lambda font, bold: (None, False))


class FallbackTests(unittest.TestCase):
    def test_fallback_matches_old_estimate(self):
        with fallback_only():
            ja_text._em_cache.clear()
            # 全角 10 字 × 20px = 200px → 幅 100px なら 2 行
            self.assertEqual(ja_text.estimate_lines("あ" * 10, 20, 100), 2)
            self.assertEqual(ja_text.estimate_lines("", 20, 100), 1)
            self.assertEqual(ja_text.estimate_lines("あ\nい", 20, 100), 2)
        ja_text._em_cache.clear()

    def test_latin_word_not_split(self):
        with fallback_only():
            ja_text._em_cache.clear()
            lines = ja_text.wrap_lines("あいう deployment", 20, 150)
            self.assertIn("deployment", lines[-1])
        ja_text._em_cache.clear()

    def test_no_line_start_hangs(self):
        with fallback_only():
            ja_text._em_cache.clear()
            lines = ja_text.wrap_lines("あいうえ。", 20, 80)
            self.assertEqual(lines, ["あいうえ。"])
        ja_text._em_cache.clear()


class BalanceTests(unittest.TestCase):
    def setUp(self):
        self.p = fallback_only()
        self.p.start()
        ja_text._em_cache.clear()

    def tearDown(self):
        self.p.stop()
        ja_text._em_cache.clear()

    def test_two_line_orphan_is_balanced(self):
        text = "四半期の売上目標を達成した要因を分析する"   # 20 字
        out = ja_text.balance(text, 20, 380)            # 19 字/行 → 最終行 1 字
        self.assertIn("\n", out)
        head, tail = out.split("\n")
        self.assertEqual(head + tail, text)
        self.assertGreaterEqual(len(tail), 3)

    def test_three_line_orphan_moves_last_break(self):
        text = "あ" * 3 + "の" + "い" * 30 + "を達成"
        out = ja_text.balance(text, 20, 300)
        self.assertEqual(out.replace("\n", ""), text)

    def test_untouched_cases(self):
        self.assertEqual(ja_text.balance("短い文", 20, 400), "短い文")
        self.assertEqual(ja_text.balance("改行\n済み" * 10, 20, 100), "改行\n済み" * 10)
        self.assertEqual(ja_text.balance("English only sentence here", 20, 100), "English only sentence here")

    def test_good_break_rules(self):
        self.assertFalse(ja_text.good_break("問い合わせ", 3))   # 送り仮名の途中
        self.assertFalse(ja_text.good_break("120分", 3))        # 数と単位
        self.assertFalse(ja_text.good_break("工数を", 2))        # 助詞は前に付ける
        self.assertTrue(ja_text.good_break("工数を削減", 3))      # 助詞の後
        self.assertFalse(ja_text.good_break("あ、", 1))          # 行頭禁則


class BuildIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.p = fallback_only()
        self.p.start()
        ja_text._em_cache.clear()

    def tearDown(self):
        self.p.stop()
        ja_text._em_cache.clear()

    def _deck(self, **meta):
        layout = b.load_json(b.ROOT / "templates" / "layouts" / "default.json")
        st = b.resolve_style(layout, {"meta": {}}, {"type": "bullets"})
        per_line = int(st["lead"]["w"] // st["lead"]["size"])
        lead = "あ" * (per_line - 6) + "の" + "い" * 7   # 自然に折り返すと最終行が 2 字になる長さ
        deck = {"meta": {"title": "t", **meta}, "slides": [{"type": "bullets", "title": "t", "lead": lead, "bullets": ["x"]}]}
        return deck, layout, lead

    def test_apply_line_balance(self):
        deck, layout, lead = self._deck()
        out = b.apply_line_balance(deck, layout)
        self.assertIn("\n", out["slides"][0]["lead"])
        self.assertEqual(out["slides"][0]["lead"].replace("\n", ""), lead)
        self.assertNotIn("\n", deck["slides"][0]["lead"])   # 元のデッキは書き換えない

    def test_line_balance_can_be_disabled(self):
        deck, layout, lead = self._deck(line_balance=False)
        self.assertEqual(b.apply_line_balance(deck, layout)["slides"][0]["lead"], lead)
        deck, layout, lead = self._deck()
        deck["slides"][0]["line_balance"] = False
        self.assertEqual(b.apply_line_balance(deck, layout)["slides"][0]["lead"], lead)


class PptxLineBreakTests(unittest.TestCase):
    def test_newline_becomes_a_br(self):
        import tempfile
        import json
        from pptx import Presentation
        deck = {"meta": {"title": "t"}, "slides": [{"type": "bullets", "title": "一行目\n二行目", "bullets": ["項目\n続き"]}]}
        with tempfile.TemporaryDirectory() as d:
            Path(d, "deck.json").write_text(json.dumps(deck, ensure_ascii=False), encoding="utf-8")
            dd = Path(d)
            deck_, theme, layout = b.load_deck(dd)
            b.build_pptx(b.expand_slides(deck_, layout), theme, layout, dd, dd / "o.pptx")
            prs = Presentation(str(dd / "o.pptx"))
            xml = "".join(sp._element.xml for sp in prs.slides[0].shapes if sp.has_text_frame)
            self.assertNotIn("一行目\n二行目", xml)
            self.assertIn("<a:br>", xml.replace("<a:br/>", "<a:br>"))
            texts = [sp.text_frame.text for sp in prs.slides[0].shapes if sp.has_text_frame]
            self.assertTrue(any("一行目" in t and "二行目" in t for t in texts))
            self.assertFalse(prs.slides[0].has_notes_slide)


if __name__ == "__main__":
    unittest.main()
